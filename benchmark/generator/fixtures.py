"""seeded 文件夹具（M8 · T1.1）：docx/xlsx/pdf 各 5 份，内嵌已知干净 PII。

确定性契约（M6 导出验收的前置）：
- docx（python-docx）/xlsx（openpyxl）保存后重写 zip（固定条目时间戳 + 固定
  core.xml 修改时间），字节级确定；
- PDF 经 ``config/synthetic_corpus.yaml`` 配置的库模块动态加载写入（固定元数据 +
  no-new-id），文本行按**槽位边界**换行打包，seeded 值绝不被换行截断，
  文本层重抽取可整串命中；
- 同 seed 两次构建 → 全部文件与 manifest 字节级一致。

夹具内容 = 模板产出的干净文书（无扰动）；manifest 记录每份夹具的 seeded PII
（类别/精确值/出现次数），供 M6 体检命中与导出零残留断言。
"""
from __future__ import annotations

import hashlib
import importlib
import random
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from benchmark.generator import templates as tpl
from benchmark.generator.personas import Personas, load_corpus_config

#: 夹具随机源 salt（与 rule_cases 的随机流隔离）
FIXTURE_SEED_SALT = 0x46495800

_PDF_PAGE_W, _PDF_PAGE_H = 595.0, 842.0
_PDF_MARGIN_X, _PDF_LINE_H = 72.0, 16.0
_PDF_MAX_COLS = 40            # 一行最大字符数（槽位整段保留，常量块按此宽切断）

_RX_MODIFIED = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def rezip_deterministic(path: Path) -> None:
    """重写 zip（原位）：固定条目时间戳/压缩级别；core.xml 修改时间固定。"""
    with zipfile.ZipFile(path, "r") as zin:
        payload = [(i.filename, i.external_attr, zin.read(i.filename))
                   for i in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=8) as zout:
        for name, attr, blob in payload:
            if name.endswith("docProps/core.xml"):
                blob = _RX_MODIFIED.sub(rb"\g<1>2026-09-28T00:00:00Z\g<2>", blob)
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = attr
            zout.writestr(zi, blob)


@dataclass
class Draft:
    """单份夹具的中间表示。"""

    filename: str
    kind: str                    # docx | xlsx | pdf
    template: str
    lines: list[str] = field(default_factory=list)            # 逻辑行
    chunk_rows: list[list[tuple[str, bool]]] = field(default_factory=list)  # 行内片段（槽位保护）
    cells: list[list[str]] = field(default_factory=list)      # xlsx：含表头
    table_rows: list[list[str]] = field(default_factory=list)  # docx：附加表格（含表头）
    seeded: list[dict[str, str]] = field(default_factory=list)
    hidden_col: str | None = None

    def text_all(self) -> str:
        rows = ["\t".join(r) for r in self.cells] + ["\t".join(r) for r in self.table_rows]
        return "\n".join([*self.lines, *rows])

    def add_seed(self, slot: tpl.Slot) -> None:
        for entry in self.seeded:
            if entry["type"] == slot.etype and entry["value"] == slot.value:
                entry["count"] = str(int(entry["count"]) + 1)  # type: ignore[union-attr]
                return
        self.seeded.append({"type": slot.etype, "value": slot.value, "count": "1"})

    def add_slots(self, slots: list[tpl.Slot | None]) -> None:
        for slot in slots:
            if slot is not None:
                self.add_seed(slot)


def _logical_chunks(chunks: list[tpl.Chunk]) -> list[list[tuple[str, bool]]]:
    """chunk 序列 → 逻辑行；每行内保留 (片段, 是否槽位值)，供 PDF 打包保护槽位。"""
    lines: list[list[tuple[str, bool]]] = []
    buf: list[tuple[str, bool]] = []
    for text, slot in chunks:
        if text == "\n":
            lines.append(buf)
            buf = []
        else:
            buf.append((text, slot is not None))
    if buf:
        lines.append(buf)
    return [ln for ln in lines if ln]


def _draft_from_template(filename: str, kind: str, template: str,
                         kit: tpl.GenKit) -> Draft:
    chunks = tpl.DOC_TEMPLATES[template](kit)
    rows = _logical_chunks(chunks)
    draft = Draft(filename=filename, kind=kind, template=template,
                  lines=["".join(t for t, _s in row) for row in rows],
                  chunk_rows=rows)
    draft.add_slots([slot for _t, slot in chunks])
    return draft


# ── 15 份夹具的装配 ────────────────────────────────────────────────


def _draft_docx_01(kit: tpl.GenKit) -> Draft:
    draft = _draft_from_template("docx_01_low_income_publicity.docx", "docx",
                                 "low_income_publicity", kit)
    rows = [["户主", "公民身份号码", "银行账户"]]
    for _ in range(3):
        person, ident = kit.person(), kit.id_card()
        bank, _bank_name = kit.bank()
        rows.append([person.value, ident.value, bank.value])
        draft.add_slots([person, ident, bank])
    draft.table_rows = rows
    return draft


def _draft_docx_03(kit: tpl.GenKit) -> Draft:
    draft = _draft_from_template("docx_03_administrative_penalty.docx", "docx",
                                 "administrative_penalty", kit)
    rows = [["当事人", "身份证号", "车辆号牌"]]
    for _ in range(2):
        person, ident, plate = kit.person(), kit.id_card(), kit.plate()
        rows.append([person.value, ident.value, plate.value])
        draft.add_slots([person, ident, plate])
    draft.table_rows = rows
    return draft


def _draft_xlsx(name: str, headers: list[str], draw_row, *,
                hidden_col: str | None = None, n_rows: int = 4,
                template: str = "roster") -> Draft:
    def build(kit: tpl.GenKit) -> Draft:
        draft = Draft(filename=name, kind="xlsx", template=template,
                      cells=[headers])
        for _ in range(n_rows):
            slots = draw_row(kit)
            draft.cells.append([s.value if s is not None else "" for s in slots])
            draft.add_slots(slots)
        draft.hidden_col = hidden_col
        return draft

    return build  # type: ignore[return-value]


def _draw_low_income_row(kit: tpl.GenKit) -> list[tpl.Slot | None]:
    person, ident = kit.person(), kit.id_card()
    bank, _bank_name = kit.bank()
    low = kit.sensitive("低保特困", kit.rng.choice(("低保对象", "低保户")))
    return [person, ident, bank, low]


def _draw_subsidy_row(kit: tpl.GenKit) -> list[tpl.Slot | None]:
    person, ident = kit.person(), kit.id_card()
    plain = tpl.Slot("OTHER", f"{kit.rng.randrange(300, 90000)}.00")
    mobile = kit.mobile()
    return [person, ident, plain, mobile]


def _draw_contact_row(kit: tpl.GenKit) -> list[tpl.Slot | None]:
    person = kit.person()
    dept = tpl.Slot("OTHER", kit.rng.choice(("办公室", "社会救助科", "社会事务科", "规划财务科")))
    landline, email = kit.landline(), kit.email()
    return [person, dept, landline, email]


def _draw_enterprise_row(kit: tpl.GenKit) -> list[tpl.Slot | None]:
    company = tpl.Slot("ORG_INTERNAL", kit.personas.company_name(), rule_detectable=False)
    uscc, legal = kit.uscc(), kit.person()
    return [company, uscc, legal, kit.address()]


def _draw_correction_row(kit: tpl.GenKit) -> list[tpl.Slot | None]:
    person, ident = kit.person(), kit.id_card()
    corr = kit.sensitive("社区矫正", "社区矫正对象")
    landline = kit.landline()
    return [person, ident, corr, landline]


XLSX_DRAFTERS = (
    _draft_xlsx("xlsx_01_low_income_roster.xlsx",
                ["户主", "公民身份号码", "银行账户", "备注"],
                _draw_low_income_row, hidden_col="D", template="low_income_publicity"),
    _draft_xlsx("xlsx_02_subsidy_payment.xlsx",
                ["姓名", "身份证号", "发放金额", "联系电话"],
                _draw_subsidy_row, template="low_income_publicity"),
    _draft_xlsx("xlsx_03_staff_contacts.xlsx",
                ["姓名", "科室", "座机", "电子邮箱"],
                _draw_contact_row, template="meeting_minutes"),
    _draft_xlsx("xlsx_04_enterprise_registry.xlsx",
                ["企业名称", "统一社会信用代码", "法定代表人", "住所"],
                _draw_enterprise_row, template="business_registration"),
    _draft_xlsx("xlsx_05_correction_roster.xlsx",
                ["姓名", "身份证号", "矫正类别", "司法所电话"],
                _draw_correction_row, template="community_correction"),
)

PDF_DRAFTERS = (
    ("pdf_01_penalty_decision.pdf", "administrative_penalty"),
    ("pdf_02_procurement_notice.pdf", "procurement_notice"),
    ("pdf_03_security_incident.pdf", "security_incident_report"),
    ("pdf_04_hotline_report.pdf", "hotline_workorder"),
    ("pdf_05_talent_publicity.pdf", "talent_publicity"),
)


def _draft_pdf(filename: str, template: str, kit: tpl.GenKit) -> Draft:
    chunks: list[tpl.Chunk] = []
    rows: list[list[tuple[str, bool]]] = []
    for _ in range(20):  # PDF 逐行渲染：密钥串若含换行（PEM 块）则重抽，确定性不变
        chunks = tpl.DOC_TEMPLATES[template](kit)
        rows = _logical_chunks(chunks)
        if not any(slot is not None and "\n" in slot.value for _t, slot in chunks):
            break
    draft = Draft(filename=filename, kind="pdf", template=template,
                  lines=["".join(t for t, _s in row) for row in rows],
                  chunk_rows=rows)
    draft.add_slots([slot for _t, slot in chunks])
    return draft


# ── 物理行打包（PDF 渲染）：槽位整段保留（最长相邻排队），常量块按 40 字切断 ──


def _pdf_physical_lines(draft: Draft) -> list[str]:
    """xlsx/docx 无换行问题（单元格/段落天然整值）；PDF 按页宽二次打包。"""
    out: list[str] = []
    for row in draft.chunk_rows:
        pieces: list[str] = []
        for text, is_slot in row:
            if is_slot:
                pieces.append(text)  # 槽位值绝不被换行截断（≤55 字符，整段成行/成段）
            else:
                pieces.extend(text[i:i + _PDF_MAX_COLS]
                              for i in range(0, len(text), _PDF_MAX_COLS))
        cur = ""
        for piece in pieces:
            if cur and len(cur) + len(piece) > _PDF_MAX_COLS:
                out.append(cur)
                cur = piece
            else:
                cur += piece
        if cur:
            out.append(cur)
    return out


# ── 三种格式写入 ───────────────────────────────────────────────────

_FIXED_META = {
    "creationDate": "D:20260928000000Z", "modDate": "D:20260928000000Z",
    "producer": "anongw-fixtures", "creator": "anongw-fixtures",
    "title": "seeded fixture", "author": "anongw", "subject": "", "keywords": "",
}


def _write_docx(draft: Draft, path: Path) -> None:
    import datetime

    from docx import Document

    doc = Document()
    doc.add_heading(draft.lines[0], level=1)
    for line in draft.lines[1:]:
        doc.add_paragraph(line)
    if draft.table_rows:
        table = doc.add_table(rows=len(draft.table_rows), cols=len(draft.table_rows[0]))
        for r, row in enumerate(draft.table_rows):
            for c, cell in enumerate(row):
                table.cell(r, c).text = cell
    props = doc.core_properties
    props.author = "anongw"
    props.created = props.modified = datetime.datetime(2026, 9, 28, 0, 0, 0)
    doc.save(path)
    rezip_deterministic(path)


def _write_xlsx(draft: Draft, path: Path) -> None:
    import datetime

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet1"
    for row in draft.cells:
        ws.append(row)
    if draft.hidden_col:
        col_idx = ws[f"{draft.hidden_col}1"].column  # 列字母 → 序号
        ws.column_dimensions[get_column_letter(col_idx)].hidden = True
    wb.properties.creator = "anongw"
    wb.properties.lastModifiedBy = "anongw"
    wb.properties.created = datetime.datetime(2026, 9, 28, 0, 0, 0)
    wb.save(path)
    rezip_deterministic(path)


def _write_pdf(draft: Draft, path: Path, config_path: str | None) -> None:
    pdfmod = importlib.import_module(
        str(load_corpus_config(config_path)["pdf_writer"]["module"]))
    doc = pdfmod.open()
    page = doc.new_page(width=_PDF_PAGE_W, height=_PDF_PAGE_H)
    y = 72.0
    for line in _pdf_physical_lines(draft):
        if y > _PDF_PAGE_H - 60:
            page = doc.new_page(width=_PDF_PAGE_W, height=_PDF_PAGE_H)
            y = 72.0
        if line:
            page.insert_text((_PDF_MARGIN_X, y), line, fontname="china-s", fontsize=10)
        y += _PDF_LINE_H
    doc.set_metadata(_FIXED_META)
    doc.save(path, deflate=True, no_new_id=True)
    doc.close()


# ── 入口 ───────────────────────────────────────────────────────────


def build_fixtures(seed: int, outdir: Path, config_path: str | None = None) -> list[dict]:
    """构建 15 份夹具并写盘，返回 manifest 条目（含 sha256 与 seeded PII）。"""
    rng = random.Random(seed)
    kit = tpl.GenKit(rng, Personas(seed, config_path))
    outdir.mkdir(parents=True, exist_ok=True)
    entries: list[dict] = []

    docx_drafts = [_draft_docx_01(kit),
                   _draft_from_template("docx_02_petition_transfer.docx", "docx",
                                        "petition_transfer", kit),
                   _draft_docx_03(kit),
                   _draft_from_template("docx_04_meeting_minutes.docx", "docx",
                                        "meeting_minutes", kit),
                   _draft_from_template("docx_05_cadre_assessment.docx", "docx",
                                        "cadre_assessment", kit)]
    xlsx_drafts = [drafter(kit) for drafter in XLSX_DRAFTERS]
    pdf_drafts = [_draft_pdf(name, template, kit) for name, template in PDF_DRAFTERS]

    for draft in (*docx_drafts, *xlsx_drafts, *pdf_drafts):
        path = outdir / draft.filename
        if draft.kind == "docx":
            _write_docx(draft, path)
        elif draft.kind == "xlsx":
            _write_xlsx(draft, path)
        else:
            _write_pdf(draft, path, config_path)
        for entry in draft.seeded:
            entry["count"] = str(draft.text_all().count(entry["value"]))
        entries.append({
            "filename": f"fixtures/{draft.filename}",
            "kind": draft.kind,
            "template": draft.template,
            "sha256": _sha256(path.read_bytes()),
            "hidden_col": draft.hidden_col,
            "seeded": draft.seeded,
        })
    return entries
