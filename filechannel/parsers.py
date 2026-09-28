"""文本层解析（M6 · T3.1）：docx / xlsx / pdf 文本层 → 带物理定位的文本段。

解析面（开发指令 §6 M6；本模块=文本层，scan_pdf 的 OCR 路径 D4 接入）：

- **docx**（python-docx，可直接 import——非禁用词表成员）：
  正文段落 + 正文表格（含嵌套表格、合并单元格去重）+ 各节页眉/页脚（含其中的
  段落与表格；``is_linked_to_previous`` 的继承节跳过，避免同文重复入段）；
- **xlsx**（openpyxl，可直接 import）：全部工作表（含隐藏表）的全部单元格——
  可见列/行与**隐藏列/隐藏行**同权重读取（隐藏正是最易漏检面，§6 M6 显式点名），
  字面值 + 公式文本（data_only=False 读取，公式串本身入段）+ 公式缓存值
  （data_only=True 读取，仅在与字面值不同时补一段）+ 单元格批注；
- **pdf 文本层**（模块名经 ``config/filechannel.yaml`` 配置、importlib 动态加载，
  铁律 E）：逐页坐标级文本抽取（textpage 全文 + 逐字符 charbox），供体检报告
  按命中 span 回填 bbox；整页零文本 → ``kind="scan_pdf"`` → **OCR 路径**
  （:mod:`filechannel.ocr`：渲染 300dpi 位图 → OCR 引擎 → 重建可检测文本 +
  命中框，T4.1 接入）。

统一产物：:class:`Segment` 列表。``location`` 形如
（:class:`filechannel.models.FileFinding` 文档串约定）：

- docx：``para:5`` / ``table:2:r1:c3``（嵌套 ``…:table:1:r1:c1``）/
  ``header:1:para:2`` / ``footer:1:table:1:r1:c1``（段落/表序号均 1 起始）；
- xlsx：``Sheet1!B3`` / ``Sheet1!B3#comment``（批注）/ ``Sheet1!C2#cached``
  （公式缓存值与公式文本同格并存时的补段）；
- pdf/scan_pdf：location 恒 ``None``——定位由 ``page`` + ``bbox`` 承担
  （§5.5 原生形状；scan_pdf 的 bbox 来自 OCR 重建层，:mod:`filechannel.ocr`）。

与 :mod:`filechannel.parse`（T3.3 导出链路的引擎面子集）并存：本模块是 T3.1
文本层权威实现（docx 全位面 / xlsx 缓存值 / pdf 坐标级 charbox）；两条链路的
收敛方向由编排层决策，本模块不依赖也不改动导出链路文件。

bbox 口径：**render 坐标系**（与 :mod:`filechannel.models` FileFinding 契约注记
一致）——原点左上、y 轴向下、单位 pt，``[x0, y0, x1, y1]`` 且 ``x0<x1、y0<y1``；
span 的 bbox = 其覆盖字符 charbox（读取库原生为左下原点）并集后翻转到 render 系。
"""
from __future__ import annotations

import importlib
import io
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import TYPE_CHECKING

import yaml

from common.config import REPO_ROOT
from filechannel.errors import DocumentParseError, UnsupportedFileType
from filechannel.models import FileKind

if TYPE_CHECKING:
    from filechannel.ocr import OcrTextLayer

#: pdf 读取库坐标（铁律 E：库名只进 config/，源码动态加载）
FILECHANNEL_CONFIG_PATH = REPO_ROOT / "config" / "filechannel.yaml"

#: 文本层支持面（扩展名 → kind；scan_pdf 由 pdf 解析结果动态判定）
SUPPORTED_SUFFIXES: dict[str, FileKind] = {
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".pdf": "pdf",
}


@dataclass(frozen=True)
class Segment:
    """一段可检测文本及其物理定位（一段=一个段落/单元格/批注/整页文本）。"""

    text: str
    page: int = 1
    location: str | None = None


BBox = tuple[float, float, float, float]


class PdfTextLayer:
    """逐页文本 + 逐字符 bbox（charbox）的并集提取器（pdf 专用）。

    charbox 存储保持读取库原生口径（左下原点、y 向上，``[l, b, r, t]``）；
    :meth:`union` 输出翻转叠加为 render 坐标系（原点左上、y 向下，契约口径）。

    索引错位退化：若某页返回文本长度与字符数不一致（UTF-16 代理对等），该页
    记入 ``fallback_pages``，span bbox 退化为整页矩形（宁可粗不可错）。
    """

    def __init__(self, page_texts: list[str], boxes: list[list[BBox]],
                 page_sizes: list[tuple[float, float]],
                 fallback_pages: set[int]) -> None:
        self.page_texts = page_texts
        self._boxes = boxes
        self._sizes = page_sizes
        self._fallback = fallback_pages

    def union(self, page: int, start: int, end: int) -> BBox | None:
        """1 起始页码 + 页内文本 span（Python 切片语义）→ render 系 bbox 并集。"""
        idx = page - 1
        if not (0 <= idx < len(self.page_texts)) or start >= end:
            return None
        w, h = self._sizes[idx]
        if idx in self._fallback:
            return (0.0, 0.0, w, h)
        row = self._boxes[idx]
        picked = [b for b in row[start:end] if b is not None]
        if not picked:
            return None
        x0 = min(b[0] for b in picked)
        x1 = max(b[2] for b in picked)
        y0 = h - max(b[3] for b in picked)   # 原生 top → render y0
        y1 = h - min(b[1] for b in picked)   # 原生 bottom → render y1
        return (min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1))


@dataclass
class ParsedDocument:
    """单文件的文本层解析产物。"""

    kind: FileKind
    pages: int
    segments: list[Segment] = field(default_factory=list)
    text_layer: PdfTextLayer | None = None  # 仅 pdf 提供（坐标级定位）
    ocr_layer: OcrTextLayer | None = None   # 仅 scan_pdf 提供（OCR 重建层）

    def locate(self, page: int, start: int, end: int) -> BBox | None:
        """span → render 系 bbox（扫描件走 OCR 层，文本层走 charbox 并集）。"""
        if self.ocr_layer is not None:
            box = self.ocr_layer.union(page, start, end)
            if box is not None:
                return box
        if self.text_layer is not None:
            return self.text_layer.union(page, start, end)
        return None


# ── docx ───────────────────────────────────────────────────────────


def _walk_docx_table(table: object, prefix: str, out: list[Segment]) -> None:
    """表格 → 单元格段（合并单元格去重；嵌套表递归，位置串拼接）。

    去重按底层 XML 元素身份：dict 值持有元素强引用，防止代理对象被 GC 后
    ``id()`` 复用导致不同单元格被误判为同一（lxml 代理为临时对象）。
    """
    seen: dict[int, object] = {}
    for r, row in enumerate(table.rows, 1):  # type: ignore[attr-defined]
        for c, cell in enumerate(row.cells, 1):  # type: ignore[attr-defined]
            tc = cell._tc  # type: ignore[attr-defined]
            key = id(tc)
            if key in seen:  # 合并单元格横向重复出现（同 _tc 元素）
                continue
            seen[key] = tc
            text = cell.text
            if text and text.strip():
                out.append(Segment(text, 1, f"{prefix}:r{r}:c{c}"))
            for nti, nested in enumerate(cell.tables, 1):  # type: ignore[attr-defined]
                _walk_docx_table(nested, f"{prefix}:r{r}:c{c}:table:{nti}", out)


def parse_docx(data: bytes) -> ParsedDocument:
    """docx 全位面文本段：正文段落 → 正文表格 → 各节页眉/页脚（含其中表格）。"""
    import docx

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 — 畸形输入统一转 422 语义
        raise DocumentParseError(f"docx 解析失败: {exc}") from exc
    out: list[Segment] = []
    for i, para in enumerate(doc.paragraphs, 1):
        if para.text and para.text.strip():
            out.append(Segment(para.text, 1, f"para:{i}"))
    for ti, table in enumerate(doc.tables, 1):
        _walk_docx_table(table, f"table:{ti}", out)
    for si, section in enumerate(doc.sections, 1):
        for part_name in ("header", "footer"):
            part = getattr(section, part_name)
            if part.is_linked_to_previous:
                continue  # 继承上一节的页眉/页脚同文，不重复入段
            for pi, para in enumerate(part.paragraphs, 1):
                if para.text and para.text.strip():
                    out.append(Segment(para.text, 1, f"{part_name}:{si}:para:{pi}"))
            for ti, table in enumerate(part.tables, 1):
                _walk_docx_table(table, f"{part_name}:{si}:table:{ti}", out)
    return ParsedDocument("docx", pages=1, segments=out)


# ── xlsx ───────────────────────────────────────────────────────────


def parse_xlsx(data: bytes) -> ParsedDocument:
    """xlsx 全工作表文本段：单元格值（可见+隐藏列/行）+ 公式文本 + 缓存值 + 批注。

    隐藏列/隐藏行不做任何跳过——openpyxl 读值与可见性无关，此处显式声明口径：
    隐藏面与可见面同权重进入体检（§6 M6「隐藏列/隐藏行」点名要求）。
    """
    import openpyxl

    try:
        wb_raw = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        wb_vals = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"xlsx 解析失败: {exc}") from exc
    out: list[Segment] = []
    for ws in wb_raw.worksheets:  # 含隐藏工作表
        ws_vals = wb_vals[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                loc = f"{ws.title}!{cell.coordinate}"
                text = str(cell.value)
                if text.strip():
                    out.append(Segment(text, 1, loc))
                cached = ws_vals[cell.coordinate].value  # 公式缓存值（无缓存则 None）
                if cached is not None and str(cached) != text and str(cached).strip():
                    out.append(Segment(str(cached), 1, f"{loc}#cached"))
                if cell.comment is not None and cell.comment.text \
                        and cell.comment.text.strip():
                    out.append(Segment(cell.comment.text, 1, f"{loc}#comment"))
    return ParsedDocument("xlsx", pages=1, segments=out)


# ── pdf 文本层 ─────────────────────────────────────────────────────


@lru_cache(maxsize=1)
def load_pdf_reader() -> object:
    """按 config/filechannel.yaml 动态加载 pdf 读取库（铁律 E：源码零库名）。

    文本层抽取（:func:`parse_pdf`）与扫描件 OCR 页面渲染
    （:func:`filechannel.ocr.ocr_page_lines`）共用同一读取库坐标。
    """
    cfg = yaml.safe_load(FILECHANNEL_CONFIG_PATH.read_text(encoding="utf-8")) or {}
    group = cfg.get("pdf_text_reader") or {}
    module_name = str(group.get("module", "")).strip()
    if not module_name:
        raise DocumentParseError(
            f"config missing pdf_text_reader.module: {FILECHANNEL_CONFIG_PATH}")
    try:
        return importlib.import_module(module_name)
    except ImportError as exc:  # noqa: TRY302 — 显式报配置/安装问题，不静默降级
        raise DocumentParseError(
            f"pdf text reader module not importable: {module_name} "
            f"(见 {FILECHANNEL_CONFIG_PATH.name})") from exc


#: 兼容别名（T3.1 原私有名；包内历史调用面）
_load_pdf_reader = load_pdf_reader


def parse_pdf(data: bytes) -> ParsedDocument:
    """pdf 解析：文本层逐页抽取（整页文本+逐字符 charbox）。

    整页零文本层 → ``kind=scan_pdf`` → OCR 路径（:mod:`filechannel.ocr`）：
    渲染位图 → OCR 引擎 → 重建可检测文本段 + 命中框层（T4.1）。
    """
    pdfium = load_pdf_reader()
    try:
        doc = pdfium.PdfDocument(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"pdf 解析失败: {exc}") from exc
    try:
        page_texts: list[str] = []
        boxes: list[list[BBox]] = []
        sizes: list[tuple[float, float]] = []
        fallback: set[int] = set()
        for page in doc:
            textpage = page.get_textpage()
            n = textpage.count_chars()
            text = textpage.get_text_range(0, n) if n else ""
            w, h = page.get_size()
            sizes.append((float(w), float(h)))
            page_texts.append(text)
            if len(text) == n:
                boxes.append([tuple(textpage.get_charbox(i)) for i in range(n)])
            else:  # 索引错位（代理对等）：该页 span bbox 退化为整页矩形
                boxes.append([])
                fallback.add(len(page_texts) - 1)
    finally:
        doc.close()
    layer = PdfTextLayer(page_texts, boxes, sizes, fallback)
    if any(t.strip() for t in page_texts):
        segments = [Segment(t, i + 1) for i, t in enumerate(page_texts) if t.strip()]
        return ParsedDocument("pdf", pages=len(page_texts), segments=segments,
                              text_layer=layer)
    # 扫描件：零文本层 → OCR 重建（延迟导入：文本层 pdf 零 OCR 开销、避免环导入）
    from filechannel.ocr import ocr_page_lines  # noqa: PLC0415

    ocr_layer = ocr_page_lines(data)
    if len(ocr_layer.pages) != len(page_texts):
        raise DocumentParseError(
            f"scan_pdf OCR 页数不一致: ocr={len(ocr_layer.pages)} "
            f"text-layer={len(page_texts)}")
    segments = []
    for i in range(len(ocr_layer.pages)):
        text = ocr_layer.page_text(i)
        if text.strip():
            segments.append(Segment(text, i + 1))
    return ParsedDocument("scan_pdf", pages=len(page_texts), segments=segments,
                          text_layer=layer, ocr_layer=ocr_layer)


# ── 分发 ───────────────────────────────────────────────────────────


def parse_any(filename: str, data: bytes) -> ParsedDocument:
    """按扩展名分发到对应解析器（不支持面 → UnsupportedFileType）。"""
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        return parse_docx(data)
    if suffix == ".xlsx":
        return parse_xlsx(data)
    if suffix == ".pdf":
        return parse_pdf(data)
    raise UnsupportedFileType(f"unsupported file type: {suffix or '<none>'} "
                              f"（支持面: {', '.join(sorted(SUPPORTED_SUFFIXES))}）")
