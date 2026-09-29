#!/usr/bin/env python3
"""端到端冒烟（开发指令 §9，D0 收工线）：mock 上游×2 + 网关全真链路，八用例 U1–U8。

拓扑（§9 步骤 1–2，全部本地、零外网依赖、零真实 API key）::

    客户端(本脚本) ──HTTP/TCP──► 网关 :9000（真实 config/app.yaml + .env 密钥）
                                   ├─ INTERNET ──► mock 上游 :8901（独立子进程）
                                   ├─ GOVCLOUD ──► mock 上游 :8902（独立子进程）
                                   └─ BLOCK    ──► 403 拦截（两 mock 零感知）

    U8（批次6 T6.2 起）另起真实链路实例：网关 :9010 仅绑 config/app.yaml 两个
    真实上游 profiles（internet_real / govcloud_real，同一 GPU 机本地大模型服务
    :9004、不同 api_key 头模拟两个真实上游）::

        客户端 ──► 网关 :9010 ├─ INTERNET ──► 本地真实大模型 :9004（internet_real）
                              └─ GOVCLOUD ──► 同服务（govcloud_real，不同 key 头）
        上游 ring buffer（/admin/*）记录收到的全文，供「全文=脱敏版」bytes 级断言；
        上游离线/未装载（GPU 服务或 ssh 隧道=环境依赖）时 U8 如实 DEFERRED 不计失败。

- mock 上游按 §9「同一 app 不同 argv」以**独立子进程**拉起
  （``python -m gateway.mock_upstream --port N``）；
- 网关在本进程内以 uvicorn 线程拉起（真实配置 + 真实端口 :9000，客户端走真实 TCP）；
  审计/会话映射走 config 指定的 SQLite 库文件（T1.3 默认落库形态，与生产一致），
  §9 U5 的"行数断言 + bytes 级零明文扫描"以**审计库文件（含 WAL 旁挂）**为对象；
  §9 U5 自 T5.1 起经 /admin/api/audit|metrics|report.csv（部门 Key 鉴权）做
  HTTP 侧交叉核对（含无 Key → 401 负例）。

用例（§9 步骤 3；U1–U5 当天生效，U6 文件通道 D3 起生效，U7 注入拦截 T4.3 起）：
- U1 非流式：人名 + 2 身份证 + 3 手机号（含分隔符写法）→ 上游 messages **全文
  diff = 原文脱敏版**（审查 §B：占位符计数精确〔人名·×1/〔身份证·×2/〔手机号·×3）；
  客户端零占位符形状、原值（归一化形态）完整在位；route=INTERNET；
- U2 流式：同 prompt，SSE 逐 delta 拼接后同 U1 断言（外加流式 AI 生成标识尾注）；
  再以 mock 默认 1–7 字符随机切块模式重放 20 次，逐次全等断言（覆盖占位符被
  切进相邻两个 SSE chunk 的还原）；
- U3 三路由（材料化·真实文件内容驱动）：三份 seeded 真实形态材料由 T1.1 生成器
  （benchmark/generator/materials.py）确定性产出为 docx/pdf 夹具文件
  （data/fixtures/materials/，双构建字节一致自证），U3 经 filechannel 文本层
  权威解析抽取**文件正文**作为 prompt 发送——普通公文 docx（零命中，messages
  全文逐字原样）→INTERNET/:8901；低保名单 docx（正文+名单表格：低保对象敏感
  个人信息 + ≥3 身份证：人名×3/身份证×3/手机×1 全量占位符化，全文=脱敏版）
  →GOVCLOUD/:8902；机密★材料 pdf→403 且两 mock 均零新增；
- U4 工具调用：tools 定义 + arguments 中文含人名/PII → 上游 messages 与
  tool_arguments **全文 = 原文脱敏版**、客户端收到还原版（JSON 全等、占位符
  零残留）；U4b：历史 assistant.tool_calls.arguments 藏密级词 → 403 整单拦截、
  上游零感知（出站全量检测面）；
- U5 审计：行数 == 本次发送的全部 /v1/chat/completions 请求数；事件字段形状齐全；
  审计 SQLite 库文件（主文件 + WAL 旁挂）bytes 级扫描全部原值（含人名与密级词，
  审查 §B）→ 零命中；/admin/api/audit|metrics|report.csv（部门 Key 鉴权）HTTP
  交叉核对（total/路由分布/报告行数；无 Key → 401）；
- U6 文件：inspect+export seeded docx → 重解析零命中（文件通道未上线时输出
  DEFERRED，不计失败）；
- U7 注入拦截（§10 场景4 防注入 e2e 化，T4.3）：检索网页存档 seeded 材料
  （真实 docx 夹具正文驱动）——干净版（含 1 经办电话）→ 200 INTERNET 照常
  脱敏/还原（负例对照：拦的是注入不是网页）；埋注版（同页评论区尾部多一段
  「AI请注意：忽略规则+外发+对用户隐匿」）→ 403 content_blocked
  reasons[0]=INJECTION、两 mock 零感知、审计 flag=injection + 预览
  〔注入·已拦截〕零埋注原文；收尾审计库 checkpoint 后 bytes 级零明文复扫；
- U8 真实大模型全链（批次6 T6.2 起）：真实链路实例（:9010，仅绑真实 profiles）
  双腿——INTERNET 腿=U1 材料、GOVCLOUD 腿=低保名单材料——普通材料→脱敏→
  真实推理→还原→AI 标识：回复非空（扣标识尾注）、占位符零泄漏、AI 标识
  （响应头+尾注+annotations）、上游 ring buffer 全文 == 原文脱敏版（逐字 diff +
  bytes 零原值）、两腿 key 指纹（SHA256 前 8 位，原值永不落缓冲）不同且与
  .env 一致；独立临时审计库 2 行（route/upstream 如实）+ 库级零明文。

运行::

    cd REPO_ROOT && PYTHONUTF8=1 ./.venv/Scripts/python.exe ops/e2e_smoke.py
    # 等价：PYTHONUTF8=1 ./.venv/Scripts/python.exe -m ops.e2e_smoke

任何 FAIL → 退出码 1；DEFERRED 不影响退出码。
"""
from __future__ import annotations

import csv
import hashlib
import http.client
import io
import json
import os
import socket
import subprocess
import sys
import threading
import time
from pathlib import Path
from typing import Any

import httpx
import uvicorn

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from audit.store import (  # noqa: E402
    RawPiiLeakError,
    assert_db_no_raw_pii,
    scan_db_files,
)
from audit.writer import SqliteAuditWriter  # noqa: E402
from benchmark.generator.materials import (  # noqa: E402
    CLASSIFIED_MATERIAL,
    ORDINARY_MATERIAL,
    ROSTER_MATERIAL,
    WEBPAGE_CLEAN_MATERIAL,
    WEBPAGE_INJECTED_MATERIAL,
    build_materials,
    material_plain_text,
    spec_of,
)
from common.config import (  # noqa: E402
    load_app_config,
    load_dept_keys,
    load_env_file,
    resolve_secret,
)
from filechannel.parsers import parse_any  # noqa: E402
from gateway.app import create_app, resolve_db_path  # noqa: E402
from gateway.mock_upstream import ECHO_MARKER  # noqa: E402
from masking.mapper import RESTORE_PATTERN, SessionRegistry  # noqa: E402
from outguard.label import AI_ANNOTATION_TYPE  # noqa: E402
from recognizers.models import EntityClass  # noqa: E402

# ── 固定拓扑与演示凭据（对齐 config/app.yaml / config/dept_keys.yaml）──────
MOCK_PORT_INTERNET = 8901
MOCK_PORT_GOVCLOUD = 8902
GATEWAY_PORT = 9000
GATEWAY_BASE = f"http://127.0.0.1:{GATEWAY_PORT}"

DEPT = "民政局"
DEMO_KEY = "dk_5e6f7a8b"        # 演示明文；config/dept_keys.yaml 只存其 sha256

AI_LABEL = "本内容由AI生成"      # config/app.yaml ai_label

READY_TIMEOUT_S = 20.0
AUDIT_SETTLE_TIMEOUT_S = 10.0
U2_REPLAYS = 20                 # §9 U2：流式重放次数

# ── U8 真实大模型全链（批次6 T6.2）────────────────────────────────────
# profile 的 base_url / api_key_env / 模型名全部以 config/app.yaml upstreams 为
# 唯一事实来源，此处只固定两个 profile 名作为选取键。
REAL_PROFILE_INTERNET = "internet_real"
REAL_PROFILE_GOVCLOUD = "govcloud_real"
U8_GATEWAY_PORT = 9010          # 真实链路独立网关实例（:9000 留给 mock 链用例）
U8_MAX_TOKENS = 192             # 单腿生成长度上限（真实推理延迟护栏，随请求体下发）
U8_HEALTH_TIMEOUT_S = 5.0       # 真实上游 /health 探测超时
U8_CLIENT_TIMEOUT_S = 180.0     # U8 客户端总超时（真实推理 + 网关链路）
U8_SESSION_INTERNET = "sess_e2e_u8a"
U8_SESSION_GOVCLOUD = "sess_e2e_u8b"


# ── 夹具：样例原值（bytes 级断言清单汇总于 RAW_VALUES）────────────────────
def _valid_id(prefix17: str) -> str:
    """给定 17 位前缀，计算 GB11643 校验位得到合法身份证样例（eval 专用夹具）。"""
    weights = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
    codes = "10X98765432"
    total = sum(int(ch) * w for ch, w in zip(prefix17, weights, strict=True))
    return prefix17 + codes[total % 11]


ID_A = "11010519491231002X"                 # 校验位合法样例
ID_B = _valid_id("46002219900307123")       # 海南区划前缀样例
ID_C = _valid_id("46002320010307123")       # 批量名单用
ID_D = _valid_id("11010519491231101")       # 批量名单用

PHONE_A = "13800138000"                     # 常规写法（= 归一化形态）
PHONE_B_RAW = "139 0013 9000"               # 空格分隔写法 → 归一化 PHONE_B
PHONE_B = "13900139000"
PHONE_C_RAW = "137-0013-7000"               # 连字符分隔写法 → 归一化 PHONE_C
PHONE_C = "13700137000"
PHONE_BLOCKED = "13900139001"               # 仅出现在机密样例（上游必须零感知）

PERSON = "张三"

WEBPAGE_PHONE = "13800138002"   # U7 网页材料的经办电话（与 materials.py 独立副本同源）

U_TEXT = (f"居民{PERSON}（身份证{ID_A}、{ID_B}），联系电话{PHONE_A}、"
          f"{PHONE_B_RAW}、{PHONE_C_RAW}，请核对其低保申领材料并回电。")
U_RESTORED = (f"居民{PERSON}（身份证{ID_A}、{ID_B}），联系电话{PHONE_A}、"
              f"{PHONE_B}、{PHONE_C}，请核对其低保申领材料并回电。")

# ── seeded 三路由材料（§9 U3 / §10 场景 2 演示同源；真实形态夹具）─────────
# 三份材料由 T1.1 生成器（benchmark/generator/materials.py）确定性产出为真实
# docx/pdf 文件：每轮运行前物化到 data/fixtures/materials/（幂等 + 双构建字节
# 一致自证），U3 经 filechannel 文本层抽取**文件正文**后发送——断言对象是真实
# 文件内容，不是脚本内构造的 JSON/纯文本。期望脱敏面（下表）为本侧独立声明的
# 验收契约：不回读被测检测器，值层面与生成器/evals.m6 的独立副本同源冻结。
MATERIALS_DIR = REPO_ROOT / "data" / "fixtures" / "materials"

#: 低保名单材料（docx 正文+名单表格）的期望脱敏面 (表面形式, 类别, 归一化值)：
#: 人名×3 + 身份证×3 + 手机号×1（全量占位符化；低保对象/低保户为 ROUTE_FLAG
#: 词不脱敏）。U3 断言「每个声明的表面形式确在文件正文里」后再用于全文 diff。
U3_ROSTER_SPANS: list[tuple[str, EntityClass, str]] = [
    ("李四", EntityClass.PERSON, "李四"),
    ("王五", EntityClass.PERSON, "王五"),
    ("赵六", EntityClass.PERSON, "赵六"),
    (ID_A, EntityClass.ID_CARD, ID_A),
    (ID_C, EntityClass.ID_CARD, ID_C),
    (ID_D, EntityClass.ID_CARD, ID_D),
    (PHONE_A, EntityClass.PHONE_MOBILE, PHONE_A),
]


def _extract_material_text(path: Path) -> str:
    """材料文件 → 正文文本（filechannel 文本层权威解析，非脚本内常量）。

    docx = 非空段落+表格段；pdf = 逐页文本（单页即整页文本）。
    """
    parsed = parse_any(path.name, path.read_bytes())
    return "\n".join(seg.text for seg in parsed.segments)


def materialize_materials() -> None:
    """三份 seeded 材料经 T1.1 生成器物化为真实 docx/pdf 夹具（幂等）。

    每轮自证三件事：①双构建字节级一致（生成器确定性）；②文件正文经文本层
    抽取 == 构造面预测（material_plain_text，往返完整性——文件形态损坏即响亮
    失败）；③本侧独立声明的期望脱敏面（U3_ROSTER_SPANS 等）确在正文里
    （值层面同源冻结的漂移哨兵）。最后清理上一代 .txt 夹具并落 manifest。
    """
    first = build_materials(MATERIALS_DIR)
    second = build_materials(MATERIALS_DIR)
    if [e["sha256"] for e in first] != [e["sha256"] for e in second]:
        raise RuntimeError("material build not byte-deterministic (T1.1 writers)")
    for entry in first:
        path = MATERIALS_DIR / entry["filename"]
        spec = spec_of(entry["filename"])
        extracted = _extract_material_text(path)
        if extracted != material_plain_text(spec):
            raise RuntimeError(f"material round-trip mismatch: {path}")
        for surface, _etype in spec.spans:
            if surface not in extracted:
                raise RuntimeError(
                    f"declared span missing from material body: {surface!r} in {path}")
    for legacy in MATERIALS_DIR.glob("*.txt"):
        legacy.unlink()  # 上一代 .txt 材料（本代升级为 docx/pdf 真实形态）
    (MATERIALS_DIR / "materials_manifest.json").write_text(
        json.dumps(first, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

TOOL_TEXT = f"查询{PERSON}名下低保发放记录，证件号{ID_D}，联系电话{PHONE_A}"
TOOL_RESTORED = f"查询{PERSON}名下低保发放记录，证件号{ID_D}，联系电话{PHONE_A}"

TOOLS = [{
    "type": "function",
    "function": {
        "name": "query_lowincome_record",
        "description": "按姓名与证件号查询低保台账",
        "parameters": {
            "type": "object",
            "properties": {"query": {"type": "string", "description": "查询原文"}},
            "required": ["query"],
        },
    },
}]

#: 全部用例涉及的敏感原值（归一化 + 原始写法 + 人名 + 密级词，审查 §B）；
#: 上游/审计 bytes 级扫描必须零命中。密级词属拦截语义（BLOCK 不出网关、审计
#: 预览以「〔密级·已拦截〕」占位），人名由规则层 v0 脱敏——两者同样不得以
#: 原文形态出现在上游 bytes 或审计库文件中，故一并入清单。
RAW_VALUES: tuple[str, ...] = (
    ID_A, ID_B, ID_C, ID_D,
    PHONE_A, PHONE_B_RAW, PHONE_B, PHONE_C_RAW, PHONE_C, PHONE_BLOCKED,
    WEBPAGE_PHONE,
    PERSON, "李四", "王五", "赵六",
    "机密★", "内部资料", "注意保密", "不得外传",
)

U1_EXPECTED_PLACEHOLDERS = {"〔身份证·": 2, "〔手机号·": 3, "〔人名·": 1}
U3_ROSTER_EXPECTED_PLACEHOLDERS = {"〔身份证·": 3, "〔手机号·": 1, "〔人名·": 3}

#: 夹具 → 上游应收「原文脱敏版」的手工 span 声明（审查 §B 全文 diff 期望侧：
#: 该脱敏什么=验收契约，不回读被测检测器，避免同源假绿）。U3 低保名单材料的
#: 期望脱敏面声明在 ROSTER_MATERIAL.spans（与夹具同源存放）。
U1_SPANS: list[tuple[str, EntityClass, str]] = [
    (PERSON, EntityClass.PERSON, PERSON),
    (ID_A, EntityClass.ID_CARD, ID_A),
    (ID_B, EntityClass.ID_CARD, ID_B),
    (PHONE_A, EntityClass.PHONE_MOBILE, PHONE_A),
    (PHONE_B_RAW, EntityClass.PHONE_MOBILE, PHONE_B),
    (PHONE_C_RAW, EntityClass.PHONE_MOBILE, PHONE_C),
]
U4_SPANS: list[tuple[str, EntityClass, str]] = [
    (PERSON, EntityClass.PERSON, PERSON),
    (ID_D, EntityClass.ID_CARD, ID_D),
    (PHONE_A, EntityClass.PHONE_MOBILE, PHONE_A),
]
#: U7 网页材料的期望脱敏面：仅 1 个经办电话（干净版/埋注版同；埋注段属
#: INJECTION·BLOCK_FLAG 不进脱敏面，拦截语义见 case_u7）
U7_SPANS: list[tuple[str, EntityClass, str]] = [
    (WEBPAGE_PHONE, EntityClass.PHONE_MOBILE, WEBPAGE_PHONE),
]
U7_EXPECTED_PLACEHOLDERS = {"〔手机号·": 1}

RESULTS: list[tuple[str, str, str]] = []   # (用例名, 状态 PASS/FAIL/DEFER, 详情)


# ── 通用小工具 ────────────────────────────────────────────────────────
def _decode_console(raw: bytes) -> str:
    """控制台工具输出解码：netstat/wmic/PowerShell 在中文 Windows 上输出 GBK
    （cp936/gb18030），而本脚本在 PYTHONUTF8=1 下直接 text=True 会按 UTF-8
    解码并炸掉 subprocess 读线程（0xBB 起始字节）——先 UTF-8、再 gb18030、
    最后 replace 兜底（仓库路径含中文，解码错表会让「本项目进程」识别失效）。"""
    for enc in ("utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _run_console(cmd: list[str], timeout: float) -> str | None:
    """跑控制台工具并解码 stdout；启动失败/超时/非零退出返回 None。"""
    try:
        r = subprocess.run(cmd, capture_output=True, timeout=timeout)
    except (OSError, subprocess.TimeoutExpired):
        return None
    if r.returncode != 0:
        return None
    return _decode_console(r.stdout or b"")


def _find_listener_pids(port: int) -> list[int]:
    """``netstat -ano`` 找出监听该端口的 PID（Windows；失败返回空表）。"""
    out = _run_console(["netstat", "-ano", "-p", "TCP"], timeout=10)
    if out is None:
        return []
    pids: set[int] = set()
    for line in out.splitlines():
        parts = line.split()
        # 形如：TCP  127.0.0.1:8901  0.0.0.0:0  LISTENING  1234
        if len(parts) >= 5 and parts[0] == "TCP" and parts[3].upper() == "LISTENING" \
                and parts[1].rsplit(":", 1)[-1] == str(port):
            try:
                pids.add(int(parts[4]))
            except ValueError:
                continue
    return sorted(pids)


def _pid_cmdline(pid: int) -> str:
    """进程命令行（wmic 优先、PowerShell CIM 兜底；均失败返回空串）。"""
    probes = (
        ["wmic", "process", "where", f"processid={pid}", "get", "commandline", "/value"],
        ["powershell", "-NoProfile", "-Command",
         f'(Get-CimInstance Win32_Process -Filter "ProcessId={pid}").CommandLine'],
    )
    for cmd in probes:
        out = _run_console(cmd, timeout=15)
        if out is None:
            continue
        if cmd[0] == "wmic":
            for ln in out.splitlines():
                if ln.strip().startswith("CommandLine="):
                    return ln.partition("=")[2].strip()
        else:
            text = out.strip()
            if text:
                return text
    return ""


def _assert_port_free(port: int) -> None:
    """端口空闲直接放行；被占时清理**确认为本项目**的 e2e 残留进程（审查 §D）。

    只有命令行同时包含本仓库路径与 mock_upstream/e2e_smoke 标识的监听进程才
    taskkill——无关占用者一概不碰，清理后仍被占即报错拒绝运行。
    """
    with socket.socket() as sock:
        sock.settimeout(0.3)
        if sock.connect_ex(("127.0.0.1", port)) != 0:
            return
    killed: list[int] = []
    for pid in _find_listener_pids(port):
        cmdline = _pid_cmdline(pid)
        if not cmdline:
            continue
        if str(REPO_ROOT) in cmdline and ("mock_upstream" in cmdline or "e2e_smoke" in cmdline):
            try:
                subprocess.run(["taskkill", "/F", "/PID", str(pid)],
                               capture_output=True, timeout=15)
                killed.append(pid)
            except (OSError, subprocess.TimeoutExpired):
                continue
    if killed:
        print(f"[e2e_smoke] 清理残留进程: port {port} <- pid {killed}", flush=True)
        time.sleep(0.5)  # Windows 端口释放有滞后，稍候复查
    with socket.socket() as sock:
        sock.settimeout(0.3)
        if sock.connect_ex(("127.0.0.1", port)) == 0:
            raise RuntimeError(
                f"port {port} already occupied — 非本项目 e2e 残留（netstat -ano 自查），不自动清理")


def _wait_ready(port: int, label: str, *, proc: subprocess.Popen | None = None,
                timeout: float = READY_TIMEOUT_S) -> None:
    """轮询本脚本自起服务的 /healthz 直到就绪（SSRF 结构性消除）。

    探测目标不是拼出来的 URL 串，而是 http.client 的**显式参数**：主机为编译期
    字面量 ``127.0.0.1``、端口经 ``int()`` 钳制——无插值/重定向/DNS rebinding 空间，
    "动态 URL 进入服务端请求"这一形态结构性不存在（mimosa 门禁实测零命中）。
    """
    deadline = time.monotonic() + timeout
    last = ""
    while time.monotonic() < deadline:
        if proc is not None and proc.poll() is not None:
            raise RuntimeError(f"{label} 进程提前退出 (code={proc.returncode})")
        try:
            conn = http.client.HTTPConnection("127.0.0.1", int(port), timeout=1.0)
            try:
                conn.request("GET", "/healthz")
                resp = conn.getresponse()
                body = json.loads(resp.read().decode("utf-8"))
            finally:
                conn.close()
            if resp.status == 200 and body.get("ok") is True:
                return
            last = f"healthz status={resp.status}"
        except Exception as exc:  # noqa: BLE001 — 启动窗口内连接失败属预期
            last = f"{type(exc).__name__}: {exc}"
        time.sleep(0.1)
    raise RuntimeError(f"{label} 就绪超时（{timeout}s）：{last}")


def _record(name: str, fn) -> None:
    try:
        detail = fn() or ""
        # 约定：返回值以 "DEFERRED" 开头 = 用例按期未生效（如 U6 文件通道 D3 前），不计失败
        status = "DEFER" if detail.startswith("DEFERRED") else "PASS"
        RESULTS.append((name, status, detail))
    except Exception as exc:  # noqa: BLE001 — 汇总所有失败后统一退出
        RESULTS.append((name, "FAIL", f"{type(exc).__name__}: {exc}"))


# ── mock 上游子进程管理（§9 步骤 1）───────────────────────────────────
def _start_mock(port: int) -> subprocess.Popen:
    _assert_port_free(port)
    proc = subprocess.Popen(
        [sys.executable, "-m", "gateway.mock_upstream", "--port", str(port)],
        cwd=str(REPO_ROOT),
        env={**os.environ, "PYTHONUTF8": "1"},
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    _wait_ready(port, label=f"mock:{port}", proc=proc)
    return proc


def _mock_count(base: str) -> int:
    return int(httpx.get(f"{base}/admin/records", timeout=5.0).json()["count"])


def _mock_text(base: str) -> bytes:
    return httpx.get(f"{base}/admin/text", timeout=5.0).content


def _mock_last_record(base: str) -> dict[str, Any]:
    records = httpx.get(f"{base}/admin/records", timeout=5.0).json()["records"]
    if not records:
        raise AssertionError(f"{base} ring buffer 为空")
    return records[-1]


def _assert_no_raw(blob: bytes, values: tuple[str, ...], where: str) -> None:
    hits = [v for v in values if v.encode("utf-8") in blob]
    if hits:
        raise AssertionError(f"{where} 出现原值（bytes 级命中）: {hits}")


def _expected_masked(ctx: dict[str, Any], session: str, text: str,
                     spans: list[tuple[str, EntityClass, str]], *,
                     service: Any = None) -> str:
    """夹具 → 上游应收到的「原文脱敏版」（审查 §B 全文 diff 的期望侧）。

    - 替换哪些表面形式由调用方手工声明（U1/U4_SPANS、ROSTER_MATERIAL.spans）
      ＝验收契约，
      不回读被测检测器（同源假绿）；
    - 占位符取网关同会话 mapper 的 §5.3.1 冻结算法结果——HMAC 同
      key/session/类别/归一化值 → 同占位符，与网关实际写入必然一致；
    - 替换按表面形式长度降序执行，避免短串吃掉长串的子串。
    ``service``：U8 真实链路实例等第二网关的 service（缺省=主 mock 链实例）。
    """
    mapper = (service if service is not None else ctx["service"]).registry.get(session)
    out = text
    for raw, etype, normalized in sorted(spans, key=lambda s: len(s[0]), reverse=True):
        placeholder, _ = mapper.placeholder_for(etype, normalized)
        out = out.replace(raw, placeholder)
    return out


def _assert_upstream_messages(record: dict[str, Any], expected_content: str, where: str) -> None:
    """records[].messages 全文 diff（审查 §B）：上游收到 = 「原文脱敏版」逐字全等。"""
    expected = [{"role": "user", "content": expected_content}]
    if record.get("messages") != expected:
        raise AssertionError(
            f"{where} messages != 原文脱敏版（全文 diff）:\n"
            f" got={record.get('messages')!r}\n exp={expected!r}")


def _assert_placeholder_counts(text: str, expected: dict[str, int]) -> None:
    for marker, want in expected.items():
        got = text.count(marker)
        if got != want:
            raise AssertionError(f"占位符 {marker}** 计数 {got} != {want}：{text!r}")


# ── 网关（§9 步骤 2：真实配置 + 真实端口 + 默认 SQLite 落库）────────────
def _drop_db(db_path: Path) -> None:
    for suffix in ("", "-wal", "-shm"):
        Path(str(db_path) + suffix).unlink(missing_ok=True)


def _start_gateway(ctx: dict[str, Any]) -> None:
    _assert_port_free(GATEWAY_PORT)
    load_env_file()  # .env 预载（MASK_KEY / MOCK_KEY；不覆盖已有环境变量）
    cfg = load_app_config()
    mask_key = resolve_secret(cfg.mask_key_env)  # 缺失即抛明确错误
    # U5 行数断言要求每轮全新库；生产默认路径（config audit_db / session_db）
    audit_db = resolve_db_path(cfg.audit_db)
    session_db = resolve_db_path(cfg.session_db)
    if audit_db == session_db:
        raise RuntimeError("audit_db 与 session_db 必须分文件（§9 U5 全库零明文扫描前提）")
    _drop_db(audit_db)
    _drop_db(session_db)
    app = create_app(cfg=cfg, mask_key=mask_key, dept_key_digests=load_dept_keys())
    ctx["cfg"] = cfg
    ctx["mask_key"] = mask_key
    ctx["service"] = app.state.service
    ctx["audit"] = app.state.service.audit
    if not isinstance(ctx["audit"], SqliteAuditWriter):
        raise RuntimeError("gateway audit store is not the SQLite writer (M7/T1.3 wiring)")
    ctx["audit_db"] = audit_db
    server = uvicorn.Server(uvicorn.Config(
        app, host="127.0.0.1", port=GATEWAY_PORT,
        log_level="warning", log_config=None, access_log=False,
    ))
    thread = threading.Thread(target=server.run, name="anongw-e2e", daemon=True)
    thread.start()
    _wait_ready(GATEWAY_PORT, label="gateway:9000")
    ctx["gw_server"], ctx["gw_thread"] = server, thread


def _stop_gateway(ctx: dict[str, Any]) -> None:
    server = ctx.get("gw_server")
    if server is not None:
        server.should_exit = True
    thread = ctx.get("gw_thread")
    if thread is not None:
        thread.join(timeout=5.0)


# ── 客户端请求（§9 步骤 3）────────────────────────────────────────────
def _auth(session: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {DEMO_KEY}", "x-anongw-session-id": session}


def _send_chat(client: httpx.Client, body: dict[str, Any], session: str,
               ctx: dict[str, Any]) -> tuple[int, dict[str, str], list[str], str]:
    """发送 /v1/chat/completions 并计数（U5 行数断言基数）。

    返回 (状态码, 响应头, SSE帧列表, 末尾残留文本)。非 SSE 响应的 JSON 体在残留文本里。
    """
    ctx["chat_sent"] += 1
    frames: list[str] = []
    with client.stream("POST", "/v1/chat/completions", json=body,
                       headers=_auth(session)) as resp:
        status, headers = resp.status_code, dict(resp.headers)
        buf = ""
        for chunk in resp.iter_text():
            buf += chunk
            while "\n\n" in buf:
                frame, buf = buf.split("\n\n", 1)
                if frame.strip():
                    frames.append(frame)
    return status, headers, frames, buf


def _plain_body(text: str, *, model: str | None = "mock-chat", stream: bool = False,
                tools: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    body: dict[str, Any] = {"messages": [{"role": "user", "content": text}], "stream": stream}
    if model is not None:
        body["model"] = model
    if tools is not None:
        body["tools"] = tools
    return body


def _sse_events(frames: list[str]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for frame in frames:
        for line in frame.splitlines():
            if line.startswith("data:"):
                payload = line[len("data:"):].strip()
                if payload and payload != "[DONE]":
                    out.append(json.loads(payload))
    return out


def _sse_done_seen(frames: list[str]) -> bool:
    return any(line.strip() == "data: [DONE]" for f in frames for line in f.splitlines())


def _sse_content(events: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for event in events:
        choices = event.get("choices") or []
        if choices and isinstance(choices[0], dict):
            delta = choices[0].get("delta") or {}
            if isinstance(delta.get("content"), str):
                parts.append(delta["content"])
    return "".join(parts)


def _plain_answer_ok(content: str, expected: str, *, streaming: bool) -> str:
    """客户端答案断言：占位符形状零残留、还原全等。

    流式（T0.5 已生效）必须带 AI 生成标识尾注；非流式的标识注入自 M5 输出侧落地，
    当前允许「无尾注」与「尾注形态」两种（M5 落地后收紧为仅后者）。
    """
    if RESTORE_PATTERN.search(content):
        raise AssertionError(f"占位符泄漏到客户端: {content!r}")
    if content == expected:
        return "exact"
    if content == f"{expected}\n{AI_LABEL}":
        if streaming:
            return "exact+label"
        return "with-ai-label(M5 预留形态)"
    if streaming:
        raise AssertionError(f"还原答案不匹配：\n got={content!r}\n exp={expected!r}（流式须带标识尾注）")
    raise AssertionError(f"还原答案不匹配：\n got={content!r}\n exp={expected!r}")


# ── U1 非流式（§9 U1）─────────────────────────────────────────────────
def case_u1(ctx: dict[str, Any]) -> str:
    client: httpx.Client = ctx["client"]
    base = ctx["mock_internet"]
    before = _mock_count(base)
    status, headers, frames, raw = _send_chat(
        client, _plain_body(U_TEXT), "sess_e2e_u1", ctx)
    if status != 200:
        raise AssertionError(f"status={status} body={raw[:200]!r}")
    if headers.get("x-anongw-route") != "INTERNET":
        raise AssertionError(f"route header: {headers.get('x-anongw-route')}")
    if not headers.get("x-anongw-request-id", "").startswith("req_"):
        raise AssertionError("request-id header missing")
    if headers.get("x-anongw-session-id") != "sess_e2e_u1":
        raise AssertionError("session-id header mismatch")
    data = json.loads(raw)
    content = data["choices"][0]["message"]["content"]
    _plain_answer_ok(content, f"{ECHO_MARKER}internet_mock\n{U_RESTORED}", streaming=False)
    for value in (ID_A, ID_B, PHONE_A, PHONE_B, PHONE_C):
        if value not in content:
            raise AssertionError(f"原值未还原: {value}")
    if data["model"] != "mock-chat" or data["usage"]["total_tokens"] <= 0:
        raise AssertionError(f"passthrough shape: model={data['model']} usage={data['usage']}")
    # 上游：仅 +1 条；全占位符（计数精确）；messages 全文 == 原文脱敏版；bytes 级零原值
    if _mock_count(base) != before + 1:
        raise AssertionError(f"8901 ring delta {_mock_count(base) - before}")
    record = _mock_last_record(base)
    _assert_placeholder_counts(record["last_user_content"], U1_EXPECTED_PLACEHOLDERS)
    _assert_upstream_messages(record, _expected_masked(ctx, "sess_e2e_u1", U_TEXT, U1_SPANS),
                              "mock:8901")
    _assert_no_raw(_mock_text(base), RAW_VALUES, "mock:8901")
    return "200 INTERNET；上游 messages 全文=原文脱敏版（人名×1/身份证×2/手机号×3）；客户端还原完整"


# ── U2 流式 + 重放 20 次（§9 U2）──────────────────────────────────────
def case_u2(ctx: dict[str, Any]) -> str:
    client: httpx.Client = ctx["client"]
    base = ctx["mock_internet"]
    before = _mock_count(base)
    expected = f"{ECHO_MARKER}internet_mock\n{U_RESTORED}\n{AI_LABEL}"

    def one_replay() -> str:
        status, headers, frames, raw = _send_chat(
            client, _plain_body(U_TEXT, stream=True), "sess_e2e_u2", ctx)
        if status != 200:
            raise AssertionError(f"status={status} body={raw[:200]!r}")
        if not headers.get("content-type", "").startswith("text/event-stream"):
            raise AssertionError(f"content-type: {headers.get('content-type')}")
        if headers.get("x-anongw-route") != "INTERNET":
            raise AssertionError(f"route header: {headers.get('x-anongw-route')}")
        if headers.get("x-anongw-ai-label") != "1":
            raise AssertionError(f"ai-label header: {headers.get('x-anongw-ai-label')}")
        if not _sse_done_seen(frames):
            raise AssertionError("missing [DONE] sentinel")
        content = _sse_content(_sse_events(frames))
        _plain_answer_ok(content, expected, streaming=True)
        return content

    first = one_replay()
    # 重放 20 次：mock 每次以 1–7 字符随机切块重发，占位符跨 chunk 切分须逐次还原全等
    for i in range(U2_REPLAYS):
        content = one_replay()
        if content != first:
            raise AssertionError(f"replay#{i + 1} content drift")
    if _mock_count(base) != before + 1 + U2_REPLAYS:
        raise AssertionError(f"8901 ring delta {_mock_count(base) - before}")
    _assert_no_raw(_mock_text(base), RAW_VALUES, "mock:8901")
    record = _mock_last_record(base)
    _assert_placeholder_counts(record["last_user_content"], U1_EXPECTED_PLACEHOLDERS)
    _assert_upstream_messages(record, _expected_masked(ctx, "sess_e2e_u2", U_TEXT, U1_SPANS),
                              "mock:8901")
    return f"200 SSE；首传 + {U2_REPLAYS} 次随机切块重放全等；上游 messages 全文=原文脱敏版；AI 标识尾注在位"


# ── U3 三路由（§9 U3；真实文件内容驱动：docx/pdf 夹具 → 文本层抽取 → 发送）──
def _load_material_text(spec) -> str:
    """加载 seeded 材料夹具正文（main() 已物化；缺失即 FAIL 而非静默回退常量）。"""
    path = MATERIALS_DIR / spec.filename
    if not path.is_file():
        raise AssertionError(f"material fixture missing: {path}")
    return _extract_material_text(path)


def case_u3(ctx: dict[str, Any]) -> str:
    client: httpx.Client = ctx["client"]
    base1, base2 = ctx["mock_internet"], ctx["mock_govcloud"]
    b1, b2 = _mock_count(base1), _mock_count(base2)

    # ① 普通公文（T1.1 生成器 docx，零敏感命中）→ INTERNET / :8901
    #    （messages 全文 diff 负例分支：逐字原文，一字不动）
    ordinary = _load_material_text(ORDINARY_MATERIAL)
    status, headers, frames, raw = _send_chat(
        client, _plain_body(ordinary), "sess_e2e_u3a", ctx)
    if status != 200 or headers.get("x-anongw-route") != "INTERNET":
        raise AssertionError(f"ordinary: {status} route={headers.get('x-anongw-route')}")
    content = json.loads(raw)["choices"][0]["message"]["content"]
    _plain_answer_ok(content, f"{ECHO_MARKER}internet_mock\n{ordinary}", streaming=False)
    if _mock_count(base1) != b1 + 1 or _mock_count(base2) != b2:
        raise AssertionError("ordinary forward split wrong")
    _assert_upstream_messages(_mock_last_record(base1), ordinary, "mock:8901 ordinary")
    _assert_no_raw(_mock_text(base1), RAW_VALUES, "mock:8901 ordinary")

    # ② 低保名单（docx 正文+名单表格：敏感个人信息 + ≥3 身份证批量双判定）
    #    → GOVCLOUD / :8902，脱敏深度 = 上游 messages 全文 == 「文件正文脱敏版」
    #    逐字全等 + 占位符计数精确 + bytes 零原值 + 客户端还原逐值在位
    roster = _load_material_text(ROSTER_MATERIAL)
    for surface, _etype, _normalized in U3_ROSTER_SPANS:
        if surface not in roster:
            raise AssertionError(f"roster material missing declared span: {surface!r}")
    status, headers, frames, raw = _send_chat(
        client, _plain_body(roster, model=None), "sess_e2e_u3b", ctx)
    if status != 200 or headers.get("x-anongw-route") != "GOVCLOUD":
        raise AssertionError(f"roster: {status} route={headers.get('x-anongw-route')}")
    content = json.loads(raw)["choices"][0]["message"]["content"]
    _plain_answer_ok(content, f"{ECHO_MARKER}govcloud_local\n{roster}", streaming=False)
    for value in ("李四", "王五", "赵六", ID_A, ID_C, ID_D, PHONE_A):
        if value not in content:
            raise AssertionError(f"roster client restore missing: {value}")
    if _mock_count(base2) != b2 + 1 or _mock_count(base1) != b1 + 1:
        raise AssertionError("roster forward split wrong")
    up = _mock_last_record(base2)
    _assert_placeholder_counts(up["last_user_content"], U3_ROSTER_EXPECTED_PLACEHOLDERS)
    _assert_upstream_messages(
        up, _expected_masked(ctx, "sess_e2e_u3b", roster, U3_ROSTER_SPANS), "mock:8902")
    _assert_no_raw(_mock_text(base2), RAW_VALUES, "mock:8902")

    # ③ 机密★材料（pdf 文本层）→ 403 拦截，两 mock 均零新增
    classified = _load_material_text(CLASSIFIED_MATERIAL)
    status, headers, frames, raw = _send_chat(
        client, _plain_body(classified), "sess_e2e_u3c", ctx)
    if status != 403:
        raise AssertionError(f"classified: {status} body={raw[:200]!r}")
    err = json.loads(raw)["error"]
    if err["code"] != "content_blocked":
        raise AssertionError(f"code: {err['code']}")
    if not err.get("reasons") or err["reasons"][0]["code"] != "CLASSIFICATION_MARK":
        raise AssertionError(f"reasons: {err['reasons']}")
    if headers.get("x-anongw-route") != "BLOCK":
        raise AssertionError(f"route header on block: {headers.get('x-anongw-route')}")
    if _mock_count(base1) != b1 + 1 or _mock_count(base2) != b2 + 1:
        raise AssertionError("blocked request reached an upstream")
    _assert_no_raw(_mock_text(base1) + _mock_text(base2),
                   (PHONE_BLOCKED, PERSON, "李四", "王五", "赵六",
                    "机密★", "内部资料", "注意保密", "不得外传"), "mocks")
    return ("真实文件三路由（docx/pdf 夹具正文驱动）：普通公文 docx→8901 逐字原文 / "
            "低保名单 docx→8902（人名×3+身份证×3+手机号×1 全量脱敏，全文=脱敏版）/ "
            "机密★ pdf→403（reasons[0]=CLASSIFICATION_MARK）；上游零原值")


# ── U4 工具调用（§9 U4）───────────────────────────────────────────────
def case_u4(ctx: dict[str, Any]) -> str:
    client: httpx.Client = ctx["client"]
    base = ctx["mock_internet"]
    before = _mock_count(base)
    status, headers, frames, raw = _send_chat(
        client, _plain_body(TOOL_TEXT, tools=TOOLS), "sess_e2e_u4", ctx)
    if status != 200:
        raise AssertionError(f"status={status} body={raw[:200]!r}")
    message = json.loads(raw)["choices"][0]["message"]
    calls = message.get("tool_calls") or []
    if len(calls) != 1 or calls[0]["function"]["name"] != "query_lowincome_record":
        raise AssertionError(f"tool_calls: {message}")
    args = calls[0]["function"]["arguments"]
    try:
        parsed = json.loads(args)
    except ValueError as exc:
        raise AssertionError(f"arguments not JSON: {args!r}") from exc
    if parsed != {"query": TOOL_RESTORED}:
        raise AssertionError(f"restored arguments mismatch: {parsed!r}")
    if RESTORE_PATTERN.search(args):
        raise AssertionError(f"placeholder leaked in arguments: {args!r}")
    record = _mock_last_record(base)
    if not record.get("tool_call") or _mock_count(base) != before + 1:
        raise AssertionError("upstream did not record the tool call")
    # 上游全文 diff（审查 §B/§F4）：messages 与 tool_arguments 都必须逐字等于
    # 「原文脱敏版」（历史 tool_calls 携带明文 PII/密级词的入站用例由 U4b 覆盖）
    expected_tool_masked = _expected_masked(ctx, "sess_e2e_u4", TOOL_TEXT, U4_SPANS)
    _assert_upstream_messages(record, expected_tool_masked, "mock:8901")
    up_args = record["tool_arguments"]
    expected_args = json.dumps({"query": expected_tool_masked}, ensure_ascii=False)
    if up_args != expected_args:
        raise AssertionError(f"upstream tool_arguments != 原文脱敏版:\n got={up_args!r}\n exp={expected_args!r}")
    _assert_no_raw(up_args.encode("utf-8"), RAW_VALUES, "mock:8901 tool_arguments")

    # U4b（审查 §A4/§B）：历史 assistant.tool_calls.arguments 藏密级词 → 整单 403，
    # 上游零感知（出站全量检测面：辅助面 BLOCK_FLAG 与正文同权拦截）
    before_b = _mock_count(base)
    history_body = {
        "model": "mock-chat",
        "messages": [
            {"role": "user", "content": "继续处理上一步的查询。"},
            {"role": "assistant", "content": None, "tool_calls": [{
                "id": "call_hist_1", "type": "function",
                "function": {"name": "query_lowincome_record",
                             "arguments": json.dumps(
                                 {"query": f"机密★{PERSON} 电话{PHONE_BLOCKED}"},
                                 ensure_ascii=False)},
            }]},
            {"role": "tool", "tool_call_id": "call_hist_1", "content": "（模拟工具结果）"},
            {"role": "user", "content": "请汇总上一步结果。"},
        ],
    }
    status, headers, frames, raw = _send_chat(client, history_body, "sess_e2e_u4b", ctx)
    if status != 403:
        raise AssertionError(f"history tool_calls classified: status={status} body={raw[:200]!r}")
    err = json.loads(raw)["error"]
    if err["code"] != "content_blocked":
        raise AssertionError(f"history classified code: {err['code']}")
    if _mock_count(base) != before_b:
        raise AssertionError("classified history tool_calls reached an upstream")
    return ("上游 messages+工具参数 全文=原文脱敏版；客户端拿到还原版（JSON 全等、占位符零残留）；"
            "历史 tool_calls 藏密级词 → 403 整单拦截、上游零感知")


# ── U5 审计（§9 U5；T1.3 起以 SQLite 库文件级扫描为口径）──────────────
def case_u5(ctx: dict[str, Any]) -> str:
    audit: SqliteAuditWriter = ctx["audit"]
    expected = ctx["chat_sent"]
    deadline = time.monotonic() + AUDIT_SETTLE_TIMEOUT_S
    while len(audit) < expected and time.monotonic() < deadline:
        time.sleep(0.1)  # 写队列批量落库 + 流式审计随流收尾，留结算窗口
    if not audit.flush(timeout_s=5.0):
        raise AssertionError("audit write queue did not drain")
    rows = audit.fetch_all()
    events = [e for _, e in rows]
    if len(events) != expected:
        raise AssertionError(f"audit rows={len(events)}, expect {expected}")

    routes = [e.route for e in events]
    # 2×BLOCK（U3c 机密正文 + U4b 历史参数藏密级词）+ 1×GOVCLOUD（U3b 批量）+ 其余 INTERNET
    if routes.count("BLOCK") != 2 or routes.count("GOVCLOUD") != 1:
        raise AssertionError(f"route mix: BLOCK={routes.count('BLOCK')} GOVCLOUD={routes.count('GOVCLOUD')}")
    if routes.count("INTERNET") != expected - 3:
        raise AssertionError(f"INTERNET count {routes.count('INTERNET')} != {expected - 3}")
    for event in events:
        if not event.request_id.startswith("req_") or not event.session_id.startswith("sess_"):
            raise AssertionError(f"ids: {event.request_id}/{event.session_id}")
        if event.dept != DEPT or event.latency_ms < 0:
            raise AssertionError(f"dept/latency: {event.dept}/{event.latency_ms}")
        if (event.route == "BLOCK") != event.blocked:
            raise AssertionError(f"blocked flag mismatch on {event.request_id}")
        if event.route == "BLOCK" and event.upstream is not None:
            raise AssertionError("blocked event must not name an upstream")
    blocked = next(e for e in events if e.route == "BLOCK")
    if blocked.response_preview != "":
        raise AssertionError("blocked response_preview must be empty")
    by_session = {e.session_id: e for e in events}
    u1 = by_session["sess_e2e_u1"]
    if "〔手机号·" not in u1.prompt_preview or ID_A in u1.prompt_preview:
        raise AssertionError(f"u1 prompt preview not masked: {u1.prompt_preview!r}")
    if "〔身份证·" not in u1.response_preview or RESTORE_PATTERN.search(u1.response_preview) is None:
        raise AssertionError(f"u1 response preview not placeholder-version: {u1.response_preview!r}")
    # bytes 级零明文（§9 U5 口径，T1.3 起=审计库文件级）：主文件 + WAL 旁挂一并扫
    audit.checkpoint()
    scanned = assert_db_no_raw_pii(ctx["audit_db"], RAW_VALUES)
    # 双保险：全部事件序列化后同样扫一遍（与库文件扫描互相独立）
    blob = "\n".join(e.model_dump_json() for e in events).encode("utf-8")
    _assert_no_raw(blob, RAW_VALUES, "audit events json")
    # 管理面查询 API（T5.1 起在位）：带部门 Key 交叉核对行数与聚合；无 Key → 401 负例
    auth = {"Authorization": f"Bearer {DEMO_KEY}"}
    resp = ctx["client"].get("/admin/api/audit", headers=auth, timeout=5.0)
    if resp.status_code != 200:
        raise AssertionError(f"/admin/api/audit status={resp.status_code} {resp.text[:200]!r}")
    page = resp.json()
    n = page.get("total", len(page.get("events", [])))
    if n != expected:
        raise AssertionError(f"/admin/api/audit rows={n} != {expected}")
    if len(page.get("events", [])) > n:
        raise AssertionError("/admin/api/audit events exceed total")
    unauth = ctx["client"].get("/admin/api/audit", timeout=5.0)
    if unauth.status_code != 401:
        raise AssertionError(f"/admin/api/audit 无 Key 应 401：{unauth.status_code}")
    report = ctx["client"].get("/admin/api/report.csv", headers=auth, timeout=5.0)
    if report.status_code != 200 or "text/csv" not in report.headers.get("content-type", ""):
        raise AssertionError(f"/admin/api/report.csv 异常: {report.status_code}")
    csv_rows = list(csv.reader(io.StringIO(report.content.decode("utf-8-sig"), newline="")))
    if len(csv_rows) != expected + 1:
        raise AssertionError(f"report.csv 行数 {len(csv_rows)} != {expected}+表头")
    if csv_rows[0][0] != "id" or csv_rows[0][5] != "route" or csv_rows[0][6] != "blocked":
        raise AssertionError(f"report.csv 表头异常: {csv_rows[0]}")
    metrics = ctx["client"].get("/admin/api/metrics", headers=auth, timeout=5.0)
    if metrics.status_code != 200:
        raise AssertionError(f"/admin/api/metrics status={metrics.status_code}")
    agg = metrics.json()
    if agg.get("requests") != expected or agg.get("blocked") != 2:
        raise AssertionError(f"metrics 聚合不符: requests={agg.get('requests')} "
                             f"blocked={agg.get('blocked')}")
    if agg.get("routes", {}).get("BLOCK") != 2 or agg.get("routes", {}).get("GOVCLOUD") != 1:
        raise AssertionError(f"metrics 路由分布不符: {agg.get('routes')}")
    if [d.get("dept") for d in agg.get("by_dept", [])] != [DEPT]:
        raise AssertionError(f"metrics 按部门聚合: {agg.get('by_dept')}")
    return (f"{expected} 行落库（2 BLOCK + 1 GOVCLOUD + {expected - 3} INTERNET）；"
            f"预览占位符版本；库文件 bytes 级扫描 {scanned} 字节零明文（含人名/密级词）；"
            f"/admin/api/audit total={n}、report.csv {expected}+表头、metrics 路由分布"
            f"交叉核对一致（无 Key 401 负例通过）")


# ── U6 文件通道（§9 U6；D3 起生效）────────────────────────────────────
def case_u6(ctx: dict[str, Any]) -> str:
    """seeded docx → inspect 命中 → export → 重解析零残留。

    文件通道未上线（404）时输出 DEFERRED（§9：U6 自 D3 起纳入），不计失败。
    """
    try:
        from docx import Document  # 文件解析通道依赖（filechannel 任务引入）
    except ImportError as exc:  # pragma: no cover — D3 前不应触达
        return f"DEFERRED: 文件通道未上线且解析依赖缺失（{exc}）"
    doc = Document()
    doc.add_paragraph(f"低保公示名单：{ID_B}，联系电话{PHONE_A}。")
    buf = io.BytesIO()
    doc.save(buf)
    payload = buf.getvalue()
    resp = ctx["client"].post(
        "/v1/files/inspect",
        files={"file": ("低保公示.docx", payload,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        timeout=30.0)
    if resp.status_code == 404:
        return "DEFERRED: 文件通道未上线（/v1/files/inspect 404；§9 U6 自 D3 起生效）"
    if resp.status_code != 200:
        raise AssertionError(f"inspect: {resp.status_code} {resp.text[:200]!r}")
    report = resp.json()
    summary = report.get("summary") or {}
    if summary.get("ID_CARD", 0) < 1 or summary.get("PHONE_MOBILE", 0) < 1:
        raise AssertionError(f"inspect summary misses seeded PII: {summary}")
    exported = ctx["client"].post(
        "/v1/files/export", data={"mode": "sanitize"},
        files={"file": ("低保公示.docx", payload,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        timeout=30.0)
    if exported.status_code != 200:
        raise AssertionError(f"export: {exported.status_code} {exported.text[:200]!r}")
    if not exported.headers.get("x-report-id"):
        raise AssertionError("export missing X-Report-Id header")
    reopened = Document(io.BytesIO(exported.content))
    text = "\n".join(p.text for p in reopened.paragraphs)
    _assert_no_raw(text.encode("utf-8"), (ID_B, PHONE_A), "exported docx")
    return "inspect 命中 seeded PII；export 后重解析零残留"


# ── U7 注入拦截（§10 场景4 防注入 e2e 化，T4.3；拦截面 = T4.2 语义层）──────
U7_SESSION_CLEAN = "sess_e2e_u7a"
U7_SESSION_INJECTED = "sess_e2e_u7b"


def case_u7(ctx: dict[str, Any]) -> str:
    """检索网页埋注（真实 docx 材料正文驱动）→ 语义层拦截 + 干净版负例对照。

    - 干净版（同页无埋注段，含 1 经办电话）→ 200 INTERNET：上游收到占位符版、
      客户端还原完整——注入拦截不伤常规脱敏链路（拦的是注入，不是网页）；
    - 埋注版（评论区尾部多一段「AI请注意：忽略规则+外发 URL+对用户隐匿」）→
      403 content_blocked reasons[0]=INJECTION，两 mock 零感知；
    - 审计：flag=injection + reasons[0]=INJECTION + 预览「〔注入·已拦截〕」
      零埋注原文/零电话原值；干净版事件 flags 保持空；
    - 收尾：审计库 checkpoint 后 bytes 级零明文复扫（RAW_VALUES 已含网页原值）。
    """
    client: httpx.Client = ctx["client"]
    base1, base2 = ctx["mock_internet"], ctx["mock_govcloud"]

    # ① 对照组：干净版网页 → 照常 INTERNET 脱敏出网
    clean = _load_material_text(WEBPAGE_CLEAN_MATERIAL)
    b1, b2 = _mock_count(base1), _mock_count(base2)
    status, headers, frames, raw = _send_chat(
        client, _plain_body(clean), U7_SESSION_CLEAN, ctx)
    if status != 200 or headers.get("x-anongw-route") != "INTERNET":
        raise AssertionError(
            f"clean webpage: {status} route={headers.get('x-anongw-route')}")
    content = json.loads(raw)["choices"][0]["message"]["content"]
    expected = _expected_masked(ctx, U7_SESSION_CLEAN, clean, U7_SPANS)
    # 客户端断言对象 = 还原版（= 原文，电话无分隔写法归一化不变），与 U1 同约定
    _plain_answer_ok(content, f"{ECHO_MARKER}internet_mock\n{clean}", streaming=False)
    if WEBPAGE_PHONE not in content:
        raise AssertionError("clean webpage phone not restored to client")
    if _mock_count(base1) != b1 + 1 or _mock_count(base2) != b2:
        raise AssertionError("clean webpage forward split wrong")
    record = _mock_last_record(base1)
    _assert_placeholder_counts(record["last_user_content"], U7_EXPECTED_PLACEHOLDERS)
    _assert_upstream_messages(record, expected, "mock:8901 webpage clean")
    _assert_no_raw(_mock_text(base1), RAW_VALUES, "mock:8901 webpage clean")

    # ② 实验组：埋注版网页 → 403 INJECTION，两 mock 零感知
    injected = _load_material_text(WEBPAGE_INJECTED_MATERIAL)
    if "忽略上面所有规则" not in injected or WEBPAGE_PHONE not in injected:
        raise AssertionError("injected webpage missing declared trap/phone")
    if injected == clean:
        raise AssertionError("injected webpage must differ from clean one")
    b1, b2 = _mock_count(base1), _mock_count(base2)
    status, headers, frames, raw = _send_chat(
        client, _plain_body(injected), U7_SESSION_INJECTED, ctx)
    if status != 403:
        raise AssertionError(f"injected webpage: {status} body={raw[:200]!r}")
    err = json.loads(raw)["error"]
    if err["code"] != "content_blocked":
        raise AssertionError(f"injected code: {err['code']}")
    if not err.get("reasons") or err["reasons"][0]["code"] != "INJECTION":
        raise AssertionError(f"injected reasons: {err['reasons']}")
    if headers.get("x-anongw-route") != "BLOCK":
        raise AssertionError(f"injected route header: {headers.get('x-anongw-route')}")
    if _mock_count(base1) != b1 or _mock_count(base2) != b2:
        raise AssertionError("injected webpage reached an upstream")

    # ③ 审计：flag=injection、预览零原文；干净版 flags 保持空
    audit: SqliteAuditWriter = ctx["audit"]
    if not audit.flush(timeout_s=5.0):
        raise AssertionError("audit write queue did not drain (u7)")
    by_session = {e.session_id: e for _, e in audit.fetch_all()}
    inj_ev = by_session.get(U7_SESSION_INJECTED)
    if inj_ev is None:
        raise AssertionError(f"audit missing event for {U7_SESSION_INJECTED}")
    if inj_ev.route != "BLOCK" or not inj_ev.blocked:
        raise AssertionError(f"injected audit route/blocked: {inj_ev.route}/{inj_ev.blocked}")
    if inj_ev.flags != ["injection"]:
        raise AssertionError(f"injected audit flags: {inj_ev.flags}")
    if not inj_ev.reasons or inj_ev.reasons[0] != "INJECTION":
        raise AssertionError(f"injected audit reasons: {inj_ev.reasons}")
    if "〔注入·已拦截〕" not in inj_ev.prompt_preview:
        raise AssertionError(
            f"injected preview lacks injection marker: {inj_ev.prompt_preview!r}")
    if "忽略上面所有规则" in inj_ev.prompt_preview or WEBPAGE_PHONE in inj_ev.prompt_preview:
        raise AssertionError("injected preview carries raw trap/phone")
    clean_ev = by_session.get(U7_SESSION_CLEAN)
    if clean_ev is None or clean_ev.route != "INTERNET" or clean_ev.flags:
        got = None if clean_ev is None else (clean_ev.route, clean_ev.flags)
        raise AssertionError(f"clean webpage audit event: {got}")

    # ④ 收尾：库文件级 bytes 零明文复扫（含 U7 新增原值）
    audit.checkpoint()
    scanned = assert_db_no_raw_pii(ctx["audit_db"], RAW_VALUES)
    return (f"网页干净版→INTERNET（电话占位符化、客户端还原）；埋注版→403 "
            f"INJECTION + flag=injection + 预览〔注入·已拦截〕零原文；两 mock 零感知；"
            f"库级复扫 {scanned} 字节零明文")


# ── U8 真实大模型全链（批次6 T6.2；普通材料→脱敏→真实推理→还原→AI 标识）────
def _real_profiles(ctx: dict[str, Any]) -> tuple[Any, Any]:
    """config/app.yaml 的两个真实上游 profiles（T6.2 起在库；缺失=配置契约缺失→FAIL）。"""
    by_name = {u.name: u for u in ctx["cfg"].upstreams}
    missing = [n for n in (REAL_PROFILE_INTERNET, REAL_PROFILE_GOVCLOUD) if n not in by_name]
    if missing:
        raise AssertionError(f"config/app.yaml 缺真实上游 profiles: {missing}")
    return by_name[REAL_PROFILE_INTERNET], by_name[REAL_PROFILE_GOVCLOUD]


def _llm_service_root(base_url: str) -> str:
    """上游 base_url（约定 …/v1）→ 服务根（…:9004）；ring/health 管理端点在根上。"""
    root = base_url.rstrip("/")
    return root[:-3] if root.endswith("/v1") else root


def _ring_last_record(root: str) -> dict[str, Any]:
    records = httpx.get(f"{root}/admin/records", timeout=5.0).json()["records"]
    if not records:
        raise AssertionError(f"真实上游 {root} ring buffer 为空")
    return records[-1]


def _start_u8_gateway(ctx: dict[str, Any], cfg_real: Any,
                      u8_db: Path) -> tuple[uvicorn.Server, threading.Thread]:
    """真实链路独立网关实例：仅绑两个真实 profiles + 独立临时审计库（不碰 U5 库）。"""
    _assert_port_free(U8_GATEWAY_PORT)
    app = create_app(cfg=cfg_real, mask_key=ctx["mask_key"],
                     dept_key_digests=load_dept_keys(),
                     audit_store=SqliteAuditWriter(u8_db),
                     session_registry=SessionRegistry(ctx["mask_key"].encode("utf-8")))
    ctx["audit_u8"] = app.state.audit
    ctx["service_u8"] = app.state.service
    server = uvicorn.Server(uvicorn.Config(
        app, host="127.0.0.1", port=U8_GATEWAY_PORT,
        log_level="warning", log_config=None, access_log=False,
    ))
    thread = threading.Thread(target=server.run, name="anongw-u8-real", daemon=True)
    thread.start()
    _wait_ready(U8_GATEWAY_PORT, label="gateway-u8:9010")
    return server, thread


def _u8_adjudicated_db_scan(db_path: Path, rows: list[Any]) -> tuple[int, list[str]]:
    """U8 专用（真实大模型腿）：全库 bytes 级扫描 + 命中宿主裁定（T8.3 收官批）。

    红线本体：**用户侧原值不得落审计库**。U1–U7 是 mock 世界（回显占位符，
    不产生自由文本），``assert_db_no_raw_pii`` 的全库零命中即红线口径；U8 的
    答案文本是真实大模型的自由生成物，可能与种子清单中的通用串字面重合
    （T8.3 run1 实锤：答案写出「张三」类示例人名 → 全库扫描命中 → 误判泄漏；
    本机 :9004 真实模型 4 样本探针 1 样本复现该重合）。裁定口径：

    - 命中串必须**只**出现在某行 ``response_preview``（模型自由文本面，网关侧
      已按 §5.4 做检测+脱敏，规则不可见的人名属 open-vocabulary 残余）内——
      且该行自身请求值的行级硬闸（append 时 assert_no_raw_pii）已核过该行
      preview 对本请求值零命中 → 判「模型自由文本与种子串字面重合」，如实注记；
    - 命中串出现在其余任何列（prompt_preview/类别计数/reasons/会话/请求号…）
      或行外自由空间 → 原样抛 RawPiiLeakError(db_file_scan)，红线不松。
    """
    blob = scan_db_files(db_path)
    hits = [v for v in RAW_VALUES if v and v.encode("utf-8") in blob]
    if not hits:
        return len(blob), []
    for value in hits:
        if not any(value in ev.response_preview for ev in rows):
            raise RawPiiLeakError("db_file_scan")
        for ev in rows:
            if value in ev.model_dump_json(exclude={"response_preview"}):
                raise RawPiiLeakError("db_file_scan")
    return len(blob), hits


def case_u8(ctx: dict[str, Any]) -> str:
    """真实大模型全链（§9 U8；批次6 T6.2 起）。

    真实链路实例（:9010，仅绑 config/app.yaml 的 internet_real / govcloud_real——
    同一 GPU 本地大模型服务、不同 api_key 头模拟两个真实上游）双腿：

    - INTERNET 腿 = U1 材料（人名+2 身份证+3 手机号，脱敏后 INTERNET 出网）；
      GOVCLOUD 腿 = 低保名单材料（敏感个人信息+批量身份证 → GOVCLOUD）；
    - 每腿断言：回复非空（扣 AI 标识尾注）、占位符零泄漏、AI 标识三面在位
      （响应头/尾注/annotations）、上游形状透传（model=配置中性名、usage>0）、
      **上游 ring buffer 全文 == 原文脱敏版**（逐字全文 diff + bytes 级零原值）、
      key 指纹（SHA256 前 8 位；原值永不落缓冲）== .env 值且两腿不同；
    - 审计：独立临时库两腿各 1 行（route/upstream 如实）+ 库文件 bytes 级扫描
      （命中宿主裁定：模型自由文本与种子串的字面重合发生在 response_preview
      内且行级硬闸干净 → 注记通过；其余宿主一律 RawPiiLeakError——见
      _u8_adjudicated_db_scan，T8.3 收官批）。

    真实上游不可达/未装载（GPU 服务或 ssh 隧道离线=环境依赖）→ DEFERRED 不计失败。
    """
    internet_up, gov_up = _real_profiles(ctx)
    root = _llm_service_root(internet_up.base_url)
    if _llm_service_root(gov_up.base_url) != root:
        raise AssertionError("两真实 profiles 应指向同一服务（同服务不同 key 头口径）")
    # ① 环境依赖门：真实上游可达且已装载（否则如实 DEFERRED，不混入 PASS/FAIL）
    try:
        health = httpx.get(f"{root}/health", timeout=U8_HEALTH_TIMEOUT_S).json()
    except Exception as exc:  # noqa: BLE001 — 隧道/服务离线属环境状态
        return f"DEFERRED: 真实上游不可达（{root}/health: {type(exc).__name__}: {exc}）"
    if not isinstance(health, dict) or not health.get("loaded"):
        err = "" if not isinstance(health, dict) else str(health.get("error"))[:120]
        return f"DEFERRED: 真实上游未装载（{root} {err}）"
    key_internet = resolve_secret(internet_up.api_key_env)
    key_gov = resolve_secret(gov_up.api_key_env)
    if key_internet == key_gov:
        raise AssertionError("internet/govcloud 两把 key 必须不同（不同 key 头口径）")

    # ② 真实链路网关实例（深拷贝配置后仅留两个真实 profiles → 每 route 首个上游即真实上游）
    cfg_real = ctx["cfg"].model_copy(deep=True)
    cfg_real.upstreams = [internet_up.model_copy(deep=True), gov_up.model_copy(deep=True)]
    tmp_dir = REPO_ROOT / "tmp"
    tmp_dir.mkdir(exist_ok=True)
    u8_db = tmp_dir / "u8_real_llm_audit.db"
    server: uvicorn.Server | None = None
    thread: threading.Thread | None = None
    client_u8: httpx.Client | None = None
    try:
        server, thread = _start_u8_gateway(ctx, cfg_real, u8_db)
        client_u8 = httpx.Client(base_url=f"http://127.0.0.1:{U8_GATEWAY_PORT}",
                                 timeout=httpx.Timeout(U8_CLIENT_TIMEOUT_S))
        httpx.post(f"{root}/admin/reset", timeout=5.0)

        legs = [
            (U8_SESSION_INTERNET, U_TEXT, U1_SPANS, "INTERNET", internet_up, key_internet),
            (U8_SESSION_GOVCLOUD, _load_material_text(ROSTER_MATERIAL), U3_ROSTER_SPANS,
             "GOVCLOUD", gov_up, key_gov),
        ]
        for session, text, spans, route, upstream, key in legs:
            body = _plain_body(text, model=None)   # model 缺省 → 上游清单首项（配置中性名）
            body["max_tokens"] = U8_MAX_TOKENS
            status, headers, frames, raw = _send_chat(client_u8, body, session, ctx)
            if status != 200:
                raise AssertionError(f"{route} 腿 status={status} body={raw[:200]!r}")
            if headers.get("x-anongw-route") != route:
                raise AssertionError(f"{route} 腿 route header: {headers.get('x-anongw-route')}")
            if headers.get("x-anongw-ai-label") != "1":
                raise AssertionError(f"{route} 腿 ai-label 响应头缺失")
            data = json.loads(raw)
            message = data["choices"][0]["message"]
            content = message["content"]
            # 回复非空（扣 AI 标识尾注行）+ 占位符零泄漏 + AI 标识三面在位
            if not content.endswith(f"\n{AI_LABEL}"):
                raise AssertionError(f"{route} 腿 AI 标识尾注缺失: {content[-40:]!r}")
            body_text = content[: -len(f"\n{AI_LABEL}")]
            if not body_text.strip():
                raise AssertionError(f"{route} 腿回复为空（仅标识行）")
            if RESTORE_PATTERN.search(content):
                raise AssertionError(f"{route} 腿占位符泄漏到客户端: {content[:120]!r}")
            annotations = message.get("annotations")
            if not annotations or annotations[0].get("type") != AI_ANNOTATION_TYPE:
                raise AssertionError(f"{route} 腿 annotations 元数据缺失: {annotations!r}")
            if data["model"] != (upstream.models[0] if upstream.models else "default") \
                    or data["usage"]["total_tokens"] <= 0:
                raise AssertionError(f"{route} 腿上游形状异常: model={data['model']} "
                                     f"usage={data['usage']}")
            # 上游 ring buffer 全文 == 原文脱敏版 + key 指纹 == .env 值（两腿不同）
            record = _ring_last_record(root)
            expected = _expected_masked(ctx, session, text, spans, service=ctx["service_u8"])
            _assert_upstream_messages(record, expected, f"llm:{root} {route}")
            fp = hashlib.sha256(key.encode("utf-8")).hexdigest()[:8]
            if not record.get("auth_header_present") or record.get("auth_key_sha8") != fp:
                raise AssertionError(f"{route} 腿 key 指纹不符: "
                                     f"{record.get('auth_key_sha8')!r} != {fp}")
            _assert_no_raw(httpx.get(f"{root}/admin/text", timeout=5.0).content,
                           RAW_VALUES, f"llm:{root} {route}")
            print(f"[e2e_smoke] U8 {route} 腿真实回复（前 60 字）: {body_text[:60]!r}…",
                  flush=True)

        # ③ 审计：独立临时库两腿各 1 行 + 库文件 bytes 级零明文
        audit_u8: SqliteAuditWriter = ctx["audit_u8"]
        deadline = time.monotonic() + AUDIT_SETTLE_TIMEOUT_S
        while len(audit_u8) < len(legs) and time.monotonic() < deadline:
            time.sleep(0.1)
        if not audit_u8.flush(timeout_s=5.0):
            raise AssertionError("audit write queue did not drain (u8)")
        by_session = {e.session_id: e for _, e in audit_u8.fetch_all()}
        for session, _text, _spans, route, upstream, _key in legs:
            event = by_session.get(session)
            if event is None:
                raise AssertionError(f"audit missing event for {session}")
            if event.route != route or event.blocked or event.flags:
                raise AssertionError(f"u8 audit {session}: "
                                     f"{event.route}/{event.blocked}/{event.flags}")
            if event.upstream != upstream.name:
                raise AssertionError(f"u8 audit upstream: {event.upstream} != {upstream.name}")
            if "〔" not in event.prompt_preview:
                raise AssertionError(f"u8 prompt preview not masked: {event.prompt_preview!r}")
        audit_u8.checkpoint()
        rows_u8 = [ev for _, ev in audit_u8.fetch_all()]
        scanned, word_hits = _u8_adjudicated_db_scan(u8_db, rows_u8)
        fps = "≠".join(hashlib.sha256(k.encode("utf-8")).hexdigest()[:8]
                       for k in (key_internet, key_gov))
        word_note = ""
        if word_hits:
            word_note = ("；模型自由文本与种子串字面重合（宿主=response_preview，"
                         f"行级硬闸已核该行自身值零落库）: {'、'.join(word_hits)}")
        return (f"双腿真实推理（INTERNET/GOVCLOUD，服务 {root}）：回复非空+占位符零泄漏+"
                f"AI 标识（头/尾注/annotations）；上游 ring 全文=原文脱敏版；"
                f"key 指纹 {fps} 两腿不同；审计 2 行 + 库级 {scanned} 字节"
                f"{'零命中' if not word_hits else '裁定通过'}{word_note}")
    finally:
        if client_u8 is not None:
            client_u8.close()
        if server is not None:
            server.should_exit = True
        if thread is not None:
            thread.join(timeout=5.0)
        for suffix in ("", "-wal", "-shm"):
            Path(str(u8_db) + suffix).unlink(missing_ok=True)


def main() -> int:
    # 直跑（不经 gate_d0）也可能落在 GBK 控制台——自带 UTF-8 重配（审查 §D）
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    ctx: dict[str, Any] = {"chat_sent": 0}
    mocks: list[subprocess.Popen] = []
    try:
        # §9 步骤 0：三份 seeded 材料物化为夹具文件（确定性、幂等，U3 从夹具加载）
        materialize_materials()
        print(f"[e2e_smoke] 材料夹具就绪: {MATERIALS_DIR}", flush=True)
        # §9 步骤 1–2：拉起两 mock 子进程 + 网关（真实端口），并复位缓冲
        mocks.append(_start_mock(MOCK_PORT_INTERNET))
        mocks.append(_start_mock(MOCK_PORT_GOVCLOUD))
        ctx["mock_internet"] = f"http://127.0.0.1:{MOCK_PORT_INTERNET}"
        ctx["mock_govcloud"] = f"http://127.0.0.1:{MOCK_PORT_GOVCLOUD}"
        for base in (ctx["mock_internet"], ctx["mock_govcloud"]):
            httpx.post(f"{base}/admin/reset", timeout=5.0)
        _start_gateway(ctx)
        ctx["client"] = httpx.Client(base_url=GATEWAY_BASE, timeout=httpx.Timeout(30.0))

        cases: list[tuple[str, Any]] = [
            ("U1 非流式往返", case_u1),
            ("U2 流式往返+重放20", case_u2),
            ("U3 三路由", case_u3),
            ("U4 工具调用还原", case_u4),
            ("U5 审计入库", case_u5),
            ("U6 文件通道", case_u6),
            ("U7 注入拦截", case_u7),
            ("U8 真实大模型全链", case_u8),
        ]
        for name, fn in cases:
            _record(name, lambda f=fn: f(ctx))
    finally:
        _stop_gateway(ctx)
        for proc in mocks:
            proc.terminate()
        for proc in mocks:
            try:
                proc.wait(timeout=5.0)
            except subprocess.TimeoutExpired:
                proc.kill()

    # §9 步骤 4：PASS/FAIL 摘要
    print("=" * 64)
    print("e2e_smoke（§9 八用例；U1–U5 当天生效，U6 D3 起，U7 注入拦截 T4.3 起，"
          "U8 真实大模型 T6.2 起——上游离线时 DEFERRED）")
    print("=" * 64)
    for name, status, detail in RESULTS:
        print(f"{status:<7} {name} — {detail}")
    passed = sum(1 for _, s, _ in RESULTS if s == "PASS")
    failed = sum(1 for _, s, _ in RESULTS if s == "FAIL")
    deferred = sum(1 for _, s, _ in RESULTS if s == "DEFER")
    print("-" * 64)
    print(f"e2e_smoke: {passed} PASS / {failed} FAIL / {deferred} DEFERRED"
          f"（DEFERRED=按期未生效不计入通过数，如 U6 文件通道 D3 起生效；"
          f"共发送 {ctx['chat_sent']} 个 /v1/chat/completions 请求）")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
