"""M8 生成器验收（T1.1）：seed 复现 + 规模/分布 + 独立校验位复核 + seeded 夹具。

运行::

    cd REPO_ROOT && PYTHONUTF8=1 ./.venv/Scripts/python.exe -m evals.m8_generator

exit 0 = 通过。检查项：
1. 双跑复现：同 seed 两次独立构建 → manifest 与全部产物字节级一致；
   data/ 规范产物与新生成物 sha256 全等（「seed 固定可复现」落到在盘文件）；
2. 评测集规模：总数 ≥300；detect/whitelist/reject 三级计分齐备；模板类 ≥10；
   扰动种类 ≥5；白名单负例 ≥40；密级样例 ≥20；难负例 ≥20；
   结构化类别正例（含扰动变体）≥20/类：ID_CARD/USCC/BANK_CARD/PHONE_MOBILE/
   PHONE_LANDLINE/PLATE/EMAIL/IP/SECRET_KEY/DATE_BIRTH；SENSITIVE_ATTR ≥20 且
   9 子类型全覆盖；PERSON/ADDRESS ≥20（NER 层计分）；澄迈身份证 ≥20；
3. 数据质量（独立实现，与生成器交叉验证）：span 恒等 raw；graded finding 的
   normalized 按同规则独立重推导全等（DATE_BIRTH 数字版式按数字组重推 ISO，
   中文数字版式仅验 ISO 形状）；ID_CARD/USCC/BANK_CARD 校验位独立重算全过；
   难负例探测值独立复算**必须**不通过校验位/合法性；
4. seeded 夹具：docx/xlsx/pdf 各 5 份在盘且 sha256 与 manifest 全等；每份夹具
   seeded PII 从文件重抽取（docx 段落+表格 / xlsx 单元格 / pdf 文本层）全部命中；
   xlsx_01 隐藏列存在；
5. 分布打印：类别/模板/扰动/敏感属性子类型计数表（§6「各类别实体计数分布打印」）。
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import tempfile
from collections import Counter
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from benchmark.generator import DEFAULT_SEED, build  # noqa: E402
from evals import thresholds as th  # noqa: E402

RESULTS: list[tuple[bool, str]] = []
DATA_DIR = REPO_ROOT / "data"

# ── 独立实现：规范化与校验位（不 import benchmark.generator 的对应物）────
_FW = {0xFF01 + i: chr(0x21 + i) for i in range(94)}
_FW[0x3000] = " "
_SEPS = " \t\r\n-‐‑‒–—―－.．·"
_ID_WEIGHTS = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
_USCC_CHARS = "0123456789ABCDEFGHJKLMNPQRTUWXY"
_USCC_WEIGHTS = (1, 3, 9, 27, 19, 26, 16, 17, 20, 29, 25, 13, 8, 24, 10, 30, 28)

#: 召回分母口径（§6 M2 对应的结构化类别）
STRUCTURAL_CLASSES = (
    "ID_CARD", "USCC", "BANK_CARD", "PHONE_MOBILE", "PHONE_LANDLINE",
    "PLATE", "EMAIL", "IP", "SECRET_KEY", "DATE_BIRTH",
)
SENSITIVE_SUBTYPES = {"病症", "残障", "低保特困", "社区矫正", "信访人",
                      "金融账户", "行踪轨迹", "犯罪记录", "特定身份"}


def _strip_seps(s: str) -> str:
    return "".join(ch for ch in s.translate(_FW) if ch not in _SEPS)


def _canon(etype: str, raw: str) -> str:
    """按 §5.3.2 独立重推导归一化目标值（与生成器实现分开编写）。"""
    hw = raw.translate(_FW)
    if etype == "ID_CARD":
        d = _strip_seps(hw)
        return d[:-1] + "X" if d.endswith(("x", "X")) else d
    if etype == "PHONE_MOBILE":
        d = _strip_seps(hw)
        body = d[1:] if d.startswith("+") else d
        if len(body) == 13 and body.startswith("86") and body[2] == "1" and body[3] in "3456789":
            return body[2:]
        return d
    if etype == "PLATE":
        return _strip_seps(hw).upper()
    if etype in ("USCC", "BANK_CARD", "PHONE_LANDLINE"):
        return _strip_seps(hw)
    return hw.strip()


def _id_ok(v: str) -> bool:
    if len(v) != 18 or not all("0" <= c <= "9" for c in v[:17]):
        return False
    tail = 10 if v[17] in ("X", "x") else (int(v[17]) if v[17].isdigit() else -1)
    if tail < 0:
        return False
    return (sum(int(c) * w for c, w in zip(v[:17], _ID_WEIGHTS, strict=True)) + tail) % 11 == 1


def _uscc_ok(v: str) -> bool:
    if len(v) != 18 or any(c not in _USCC_CHARS for c in v):
        return False
    idx = {c: i for i, c in enumerate(_USCC_CHARS)}
    total = sum(idx[c] * w for c, w in zip(v[:17], _USCC_WEIGHTS, strict=True)) + idx[v[17]]
    return total % 31 == 0


def _luhn_ok(v: str) -> bool:
    if not v.isdigit() or len(v) < 12:
        return False
    total = 0
    for i, ch in enumerate(reversed(v)):
        d = int(ch)
        if i % 2 == 1:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def _record(name: str, fn):
    try:
        detail = fn()
        RESULTS.append((True, f"PASS  {name}" + (f" — {detail}" if detail else "")))
    except Exception as exc:  # noqa: BLE001 — 汇总所有失败后统一退出
        RESULTS.append((False, f"FAIL  {name} — {type(exc).__name__}: {exc}"))


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# ── 1. 双跑复现 + 在盘一致 ─────────────────────────────────────────


def check_seed_repro() -> str:
    with tempfile.TemporaryDirectory() as td:
        manifest_a = build.build_all(DEFAULT_SEED, Path(td) / "a")
        build.build_all(DEFAULT_SEED, Path(td) / "b")
        ta = (Path(td) / "a" / "generation_manifest.json").read_bytes()
        tb = (Path(td) / "b" / "generation_manifest.json").read_bytes()
        assert ta == tb, "two runs produced different manifests"
        mismatches = [rel for rel, sha in manifest_a["files"].items()
                      if _sha((Path(td) / "a" / rel).read_bytes()) != sha
                      or _sha((Path(td) / "b" / rel).read_bytes()) != sha]
        assert not mismatches, f"artifact bytes differ across runs: {mismatches}"

    manifest_path = DATA_DIR / "generation_manifest.json"
    assert manifest_path.exists(), "data/generation_manifest.json missing (run: python -m benchmark.generator)"
    on_disk = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert on_disk["seed"] == DEFAULT_SEED, f"data seed {on_disk['seed']} != {DEFAULT_SEED}"
    fresh = build.build_all(DEFAULT_SEED, DATA_DIR / "_m8_fresh_check")
    try:
        drifted = [rel for rel, sha in fresh["files"].items()
                   if not (DATA_DIR / rel).exists()
                   or _sha((DATA_DIR / rel).read_bytes()) != sha]
        assert not drifted, f"data/ artifacts drifted from seed {DEFAULT_SEED}: {drifted}"
    finally:
        import shutil

        shutil.rmtree(DATA_DIR / "_m8_fresh_check", ignore_errors=True)
    return (f"two runs byte-identical ({len(fresh['files'])} files); "
            f"data/ == seed {DEFAULT_SEED} rebuild")


# ── 2/3. 评测集规模 + 数据质量 ─────────────────────────────────────


def _load_cases() -> list[dict]:
    path = DATA_DIR / "rule_cases.jsonl"
    assert path.exists(), "data/rule_cases.jsonl missing"
    mirror = DATA_DIR / "cases" / "rule_cases.jsonl"
    assert mirror.exists(), "data/cases/rule_cases.jsonl mirror missing"
    raw = path.read_bytes()
    assert raw == mirror.read_bytes(), "mirror is not byte-identical"
    cases = [json.loads(line) for line in raw.decode("utf-8").splitlines() if line.strip()]
    assert len({c["id"] for c in cases}) == len(cases), "duplicate case ids"
    return cases


def check_case_scale(cases: list[dict]) -> str:
    assert len(cases) >= th.RULE_CASES_MIN, f"cases {len(cases)} < {th.RULE_CASES_MIN}"
    grading = Counter(c["grading"] for c in cases)
    assert grading.get("whitelist", 0) >= th.WHITELIST_NEGATIVES_MIN, grading
    assert grading.get("reject", 0) >= th.GENERATOR_REJECT_NEGATIVES_MIN, grading
    assert {"detect", "whitelist", "reject"} <= set(grading), grading

    templates = {c["template"] for c in cases}
    assert len(templates) >= th.GENERATOR_TEMPLATE_CLASSES_MIN, sorted(templates)
    doc_templates = {c["template"] for c in cases if not c["template"].startswith(("wl_", "r_"))}
    assert len(doc_templates) >= th.GENERATOR_TEMPLATE_CLASSES_MIN, sorted(doc_templates)

    perturbs = {k for c in cases for k in c["perturbations"]}
    assert len(perturbs) >= th.GENERATOR_PERTURB_KINDS_MIN, sorted(perturbs)

    def cases_with(etype: str, *, graded_only: bool = True) -> int:
        return sum(1 for c in cases if any(
            f["type"] == etype and (f["graded"] or not graded_only) for f in c["expected"]))

    for etype in STRUCTURAL_CLASSES:
        n = sum(1 for c in cases if any(
            f["type"] == etype and f["graded"] and f["rule_detectable"]
            and not f["whitelisted"] for f in c["expected"]))
        floor = (th.GENERATOR_DATE_BIRTH_MIN if etype == "DATE_BIRTH"
                 else th.GENERATOR_LANDLINE_MIN if etype == "PHONE_LANDLINE"
                 else th.RULE_CASES_PER_CLASS_MIN)
        assert n >= floor, f"{etype}: {n} < {floor}"
    for etype, floor in (("PERSON", th.GENERATOR_PERSON_MIN),
                         ("ADDRESS", th.GENERATOR_ADDRESS_MIN),
                         ("SENSITIVE_ATTR", th.GENERATOR_SENSITIVE_ATTR_MIN)):
        n = cases_with(etype)
        assert n >= floor, f"{etype}: {n} < {floor}"

    subtypes = {f["subtype"] for c in cases for f in c["expected"]
                if f["type"] == "SENSITIVE_ATTR" and f["subtype"]}
    assert subtypes == SENSITIVE_SUBTYPES, f"subtype coverage: {sorted(subtypes)}"

    n_class = sum(1 for c in cases
                  if any(f["type"] == "CLASSIFICATION_MARK" for f in c["expected"]))
    assert n_class >= th.CLASSIFICATION_SAMPLES_MIN, n_class
    n_cm = sum(1 for c in cases for f in c["expected"]
               if f["type"] == "ID_CARD" and f["graded"]
               and f["normalized"].startswith("469023"))
    assert n_cm >= th.GENERATOR_ID_CHENGMAI_MIN, n_cm
    return (f"{len(cases)} cases (detect={grading['detect']}, whitelist={grading['whitelist']}, "
            f"reject={grading['reject']}), {len(doc_templates)} doc templates, "
            f"{len(perturbs)} perturbation kinds, chengmai ids={n_cm}")


def check_case_quality(cases: list[dict]) -> str:
    bad_span, bad_norm, bad_chk, bad_reject = [], [], [], []
    date_seen = 0
    for case in cases:
        text = case["text"]
        for f in case["expected"]:
            if text[f["start"]:f["end"]] != f["raw"]:
                bad_span.append((case["id"], f["fid"]))
            if not f["graded"]:
                continue
            if f["type"] == "DATE_BIRTH":
                digits = re.findall(r"\d+", f["raw"])
                if len(digits) == 3:
                    date_seen += 1
                    iso = f"{int(digits[0]):04d}-{int(digits[1]):02d}-{int(digits[2]):02d}"
                    if iso != f["normalized"]:
                        bad_norm.append((case["id"], f["fid"], "date"))
                elif not re.fullmatch(r"\d{4}-\d{2}-\d{2}", f["normalized"]):
                    bad_norm.append((case["id"], f["fid"], "date-shape"))
                continue
            if _canon(f["type"], f["raw"]) != f["normalized"]:
                bad_norm.append((case["id"], f["fid"], f["type"]))
            if f["type"] == "ID_CARD" and not _id_ok(f["normalized"]):
                bad_chk.append((case["id"], f["normalized"]))
            if f["type"] == "USCC" and not _uscc_ok(f["normalized"]):
                bad_chk.append((case["id"], f["normalized"]))
            if f["type"] == "BANK_CARD" and not _luhn_ok(f["normalized"]):
                bad_chk.append((case["id"], f["normalized"]))
        probe = case.get("reject_probe")
        if probe is not None:
            v, etype = probe["raw"], probe["type"]
            ok = {"ID_CARD": _id_ok, "USCC": _uscc_ok, "BANK_CARD": _luhn_ok}.get(etype)
            invalid = (not ok(v)) if ok else (
                etype == "PHONE_MOBILE" and len(_canon(etype, v)) != 11
                or etype == "IP" and any(int(o) > 255 for o in v.split(".")))
            if not invalid:
                bad_reject.append((case["id"], etype, v))
    assert not bad_span, f"span mismatches: {bad_span[:5]}"
    assert not bad_norm, f"normalized mismatches: {bad_norm[:5]}"
    assert not bad_chk, f"checksum failures: {bad_chk[:5]}"
    assert not bad_reject, f"reject probes unexpectedly valid: {bad_reject[:5]}"
    assert date_seen >= th.GENERATOR_DATE_BIRTH_MIN, date_seen
    return (f"spans ok; normalized ok ({date_seen} digit-format dates re-derived); "
            f"ID/USCC/Luhn recomputed ok; {sum(1 for c in cases if 'reject_probe' in c)} "
            f"reject probes independently invalid")


# ── 4. seeded 夹具 ─────────────────────────────────────────────────


def check_fixtures() -> str:
    manifest = json.loads((DATA_DIR / "generation_manifest.json").read_text(encoding="utf-8"))
    fixtures = manifest["fixtures"]
    kinds = Counter(e["kind"] for e in fixtures)
    assert kinds == {"docx": th.FILE_FIXTURES_PER_KIND, "xlsx": th.FILE_FIXTURES_PER_KIND,
                     "pdf": th.FILE_FIXTURES_PER_KIND}, dict(kinds)

    import docx
    import openpyxl

    from benchmark.generator.personas import load_corpus_config

    pdfmod = __import__(load_corpus_config()["pdf_writer"]["module"])

    misses: list[str] = []
    hidden_seen = False
    for entry in fixtures:
        path = DATA_DIR / entry["filename"]
        assert path.exists(), f"missing fixture: {entry['filename']}"
        assert _sha(path.read_bytes()) == entry["sha256"], f"sha drift: {entry['filename']}"
        if entry["kind"] == "docx":
            d = docx.Document(path)
            text = "\n".join(p.text for p in d.paragraphs)
            text += "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
        elif entry["kind"] == "xlsx":
            wb = openpyxl.load_workbook(path)
            ws = wb[wb.sheetnames[0]]
            hidden = [k for k, v in ws.column_dimensions.items() if v.hidden]
            if entry.get("hidden_col"):
                assert entry["hidden_col"] in hidden, f"hidden col missing: {entry['filename']}"
                hidden_seen = True
            text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row)
        else:
            doc = pdfmod.open(path)
            text = "".join(page.get_text() for page in doc)
            doc.close()
        for seeded in entry["seeded"]:
            if int(seeded["count"]) > 0 and seeded["value"] not in text:
                misses.append(f"{entry['filename']}:{seeded['type']}:{seeded['value'][:24]}")
    assert not misses, f"seeded PII not round-tripped: {misses[:5]}"
    assert hidden_seen, "no hidden-column xlsx fixture verified"
    return (f"{len(fixtures)} fixtures (docx/xlsx/pdf ×{th.FILE_FIXTURES_PER_KIND}), "
            f"sha256 ok, seeded PII round-trip ok, hidden col ok")


# ── 5. 分布打印（§6：各类别实体计数分布）───────────────────────────


def print_distributions(cases: list[dict]) -> None:
    manifest = json.loads((DATA_DIR / "generation_manifest.json").read_text(encoding="utf-8"))
    counts = manifest["counts"]
    print("── 分布：按类别（graded-detect 正例 / 全部 finding）─────────────")
    all_f = counts["findings_by_class"]
    graded = counts["graded_detect_by_class"]
    for etype in sorted(all_f):
        print(f"  {etype:<18} graded={graded.get(etype, 0):<4} total={all_f[etype]}")
    print("── 分布：按模板 ────────────────────────────────────────────────")
    for name, n in counts["templates"].items():
        print(f"  {name:<28} {n}")
    print("── 分布：扰动 ──────────────────────────────────────────────────")
    for kind, n in counts["perturbations"].items():
        print(f"  {kind:<18} {n}")
    print("── 分布：敏感属性子类型 ────────────────────────────────────────")
    for st, n in counts["sensitive_attr_by_subtype"].items():
        print(f"  {st:<10} {n}")


def main() -> int:
    _record("seed-repro(双跑字节级+在盘一致)", check_seed_repro)
    cases = _load_cases()
    _record("case-scale(≥300/三级计分/模板/扰动/类别)", lambda: check_case_scale(cases))
    _record("case-quality(span/normalized/校验位/难负例)", lambda: check_case_quality(cases))
    _record("fixtures(5×3/sha/重抽取命中/隐藏列)", check_fixtures)
    try:
        print_distributions(cases)
    except Exception as exc:  # noqa: BLE001 — 打印失败不影响判定
        print(f"distribution print skipped: {exc}")

    for _, line in RESULTS:
        print(line)
    passed = sum(1 for ok, _ in RESULTS if ok)
    total = len(RESULTS)
    print(f"M8 GENERATOR: {passed}/{total} checks passed")
    return 0 if passed == total else 1


if __name__ == "__main__":
    sys.exit(main())
