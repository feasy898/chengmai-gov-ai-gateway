"""M6 文件通道·文本层验收（T3.1）：解析 → detect 全套 → FileReport 逐 Finding 带位置。

运行::

    cd REPO_ROOT && PYTHONUTF8=1 ./.venv/Scripts/python.exe -m evals.m6_filesvc

范围（任务单 T3.1=文本层；导出/重解析零残留属后续任务，不入本门）：
被测对象为 :mod:`filechannel.parsers`（docx 全位面 / xlsx 含隐藏列与批注 /
pdf 坐标级文本层）+ :func:`filechannel.inspect.inspect_bytes`（detect 全套 →
FileReport 装配、§5.5 风险分级）。与 T3.3 导出链路（filechannel/parse.py +
service.py，与本模块并行落地）互不依赖。

通过线（开发指令 §6 M6 文本层部分 / evals.thresholds）：

1. seeded 夹具命中率 100%：M8 生成的 15 份夹具（docx/xlsx/pdf ×
   FILE_FIXTURES_PER_KIND）逐份 inspect，每条可评级 seeded PII（类别一致 +
   归一化等价）命中数 ≥ 声明 count；总数 ≥ FILE_GRADED_SEEDED_MIN；
2. 隐藏面可见：xlsx_01 隐藏列 D 的敏感属性命中于 ``Sheet1!D*``（自造夹具另验
   隐藏列/隐藏行/批注三面的精确定位）；
3. docx 全位面：正文段落/表格/页眉/页脚逐位置命中（``para:`` ``table:``
   ``header:`` ``footer:``）；
4. pdf 坐标级：逐 Finding bbox 齐备、在页内（render 系原点左上），并与解析层
   charbox 并集复核全等、与库级「按框取文」回读文本一致；
5. 零文本层 pdf → kind=scan_pdf、零命中（OCR 路径 D4 接入前的如实降级）；
6. 三路由材料（data/fixtures/materials，与 ops/e2e_smoke 同源）材料化为
   docx/xlsx/pdf 后逐 span 命中；
7. 风险分级 §5.5 字面口径：HIGH=SENSITIVE_ATTR/批量身份证（BANK_CARD 并入批量，
   阈值同 §5.2 路由批量线）；MID=其它非白名单结构化 PII；LOW=仅白名单/词面；
   NONE=零命中（密级/工作秘密属 BLOCK/ROUTE 语义，不进分级阶梯但仍出 finding）；
8. 边界：>50MB 拒绝、不支持种类拒绝、畸形文件解析失败（一律异常语义，不崩进程）。

评级分母口径（与 m2_recognizers 同源）：PERSON/ADDRESS/ORG_INTERNAL/OTHER 归
NER 层、WORK_SECRET 属词表可选层——均不入「必须命中」分母（命中算加分）；
DATE_BIRTH 仅计数字可判版式（中文数字版式规则层不可判，不入分母）。
夹具期望来自 M8 manifest（data/generation_manifest.json，seed 固定可再生；
缺失/漂移先跑 evals.m8_generator 重建）。
"""
from __future__ import annotations

import json
import random
import re
import sys
from pathlib import Path

import yaml

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from benchmark.generator import numbers as nums  # noqa: E402
from benchmark.generator.personas import load_corpus_config  # noqa: E402
from evals import thresholds as th  # noqa: E402
from filechannel import UnsupportedFileType, inspect_bytes  # noqa: E402
from filechannel.errors import DocumentParseError, FileTooLargeError  # noqa: E402
from filechannel.inspect import risk_level_of  # noqa: E402
from filechannel.models import FileFinding, FileReport  # noqa: E402
from filechannel.parsers import parse_pdf  # noqa: E402
from masking.normalize import normalize_value  # noqa: E402
from recognizers.models import EntityClass, Finding  # noqa: E402

RESULTS: list[tuple[bool, str]] = []
DATA_DIR = REPO_ROOT / "data"
MANIFEST_PATH = DATA_DIR / "generation_manifest.json"
OUT_DIR = DATA_DIR / "fixtures" / "filesvc"
MATERIALS_OUT_DIR = OUT_DIR / "materials"

#: 不入规则层评级分母的类别（NER 层 / 词表可选层；与生成器 rule_detectable 口径一致）
NER_LAYER_TYPES = frozenset({"PERSON", "ADDRESS", "ORG_INTERNAL", "OTHER", "WORK_SECRET"})
_ID_DATE_SHAPE = re.compile(r"(?:18|19|20)\d{2}[-/.]\d{1,2}[-/.]\d{1,2}"
                            r"|(?:18|19|20)\d{2}年\d{1,2}月\d{1,2}日")


def _record(name: str, fn):
    try:
        detail = fn()
        RESULTS.append((True, f"PASS  {name}" + (f" — {detail}" if detail else "")))
    except Exception as exc:  # noqa: BLE001 — 汇总所有失败后统一退出
        RESULTS.append((False, f"FAIL  {name} — {type(exc).__name__}: {exc}"))


def _gradeable(etype: str, value: str) -> bool:
    """seeded 值是否入规则层评级分母（m2 同口径；见模块 docstring）。"""
    if etype in NER_LAYER_TYPES:
        return False
    if etype == "DATE_BIRTH":
        return bool(_ID_DATE_SHAPE.fullmatch(value))
    return True


def _matches(report: FileReport, etype: str, value: str) -> list[FileFinding]:
    """类别一致 + 归一化等价（§5.3.2：同号异写 → 同 normalized）的命中列表。"""
    want = normalize_value(EntityClass(etype), value)
    return [ff for ff in report.findings
            if ff.finding.type.value == etype and ff.finding.normalized == want]


def _load_fixtures() -> list[dict]:
    assert MANIFEST_PATH.exists(), \
        f"missing {MANIFEST_PATH}（先跑 evals.m8_generator 再生 seeded 夹具）"
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    fixtures = manifest.get("fixtures", [])
    kinds = {k: sum(1 for e in fixtures if e["kind"] == k) for k in ("docx", "xlsx", "pdf")}
    assert all(n == th.FILE_FIXTURES_PER_KIND for n in kinds.values()), kinds
    return fixtures


# ── 1. seeded 夹具全量体检 ─────────────────────────────────────────


def check_fixtures(fixtures: list[dict]) -> str:
    import hashlib

    graded_total = matched_total = 0
    location_misses: list[str] = []
    for entry in fixtures:
        path = DATA_DIR / entry["filename"]
        assert path.exists(), f"missing fixture: {entry['filename']}"
        data = path.read_bytes()
        assert hashlib.sha256(data).hexdigest() == entry["sha256"], \
            f"sha drift: {entry['filename']}（先跑 evals.m8_generator 重建）"
        report = inspect_bytes(path.name, data)
        assert report.kind == entry["kind"], (entry["filename"], report.kind)
        assert report.sha256 == entry["sha256"]
        assert report.file_id.startswith("file_") and len(report.file_id) > 5
        # fid 文件内自增连续（§5.1「请求内自增」的文件内口径）
        assert [ff.finding.fid for ff in report.findings] == \
            [f"f_{i:04d}" for i in range(1, len(report.findings) + 1)]
        assert sum(report.summary.values()) == len(report.findings)
        for ff in report.findings:
            if report.kind == "pdf":
                assert ff.page >= 1 and ff.bbox is not None and len(ff.bbox) == 4, \
                    f"pdf finding missing page/bbox: {entry['filename']}"
            else:
                assert ff.page == 1 and ff.location, \
                    f"{report.kind} finding missing location: {entry['filename']}"
        for seeded in entry["seeded"]:
            if not _gradeable(seeded["type"], seeded["value"]):
                continue
            graded_total += 1
            hits = _matches(report, seeded["type"], seeded["value"])
            if len(hits) >= int(seeded["count"]):
                matched_total += 1
            else:
                location_misses.append(
                    f"{entry['filename']}:{seeded['type']}:{seeded['value'][:18]}"
                    f" got={len(hits)} want>={seeded['count']}")
    assert not location_misses, f"seeded misses: {location_misses[:5]}"
    assert graded_total >= th.FILE_GRADED_SEEDED_MIN, graded_total
    return (f"{len(fixtures)} fixtures seeded 命中 {matched_total}/{graded_total}"
            f"（person/address 等词表外加成不计分母）")


# ── 2. 隐藏列（M8 夹具 xlsx_01 + 自造三面夹具）─────────────────────


def check_xlsx_hidden_col(fixtures: list[dict]) -> str:
    import openpyxl

    entry = next(e for e in fixtures if e.get("hidden_col"))
    path = DATA_DIR / entry["filename"]
    ws = openpyxl.load_workbook(path).worksheets[0]
    assert ws.column_dimensions[entry["hidden_col"]].hidden, "夹具自证：隐藏列未隐藏"

    report = inspect_bytes(path.name, path.read_bytes())
    col = f"{ws.title}!{entry['hidden_col']}"
    in_hidden = [ff for ff in report.findings
                 if ff.location and ff.location.startswith(col)
                 and ff.finding.type is EntityClass.SENSITIVE_ATTR]
    want = sum(int(s["count"]) for s in entry["seeded"]
               if s["type"] == "SENSITIVE_ATTR")
    assert len(in_hidden) >= want, (len(in_hidden), want)
    assert all(ff.bbox is None for ff in report.findings)  # xlsx 无 bbox
    return (f"{entry['filename']} 隐藏列 {col}* 命中 SENSITIVE_ATTR "
            f"{len(in_hidden)}/{want}")


def _selfmade_xlsx(path: Path) -> dict[str, str]:
    """自造含隐藏列/隐藏行/批注的 xlsx（任务单点名面）；返回 值→期望定位 表。"""
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    rng = random.Random(20260929)
    region, _ = nums.pick_region(rng)
    id1 = nums.gen_id_card(rng, region, nums.gen_birth(rng))
    id2 = nums.gen_id_card(rng, region, nums.gen_birth(rng))
    id3 = nums.gen_id_card(rng, region, nums.gen_birth(rng))
    uscc = nums.gen_uscc(rng, region)
    bank, _name = nums.gen_bank_card(rng)
    mobile1, mobile2 = nums.gen_mobile(rng), nums.gen_mobile(rng)
    email = nums.gen_email(rng)
    plate = nums.gen_plate(rng)

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet1"
    ws.append(["姓名", "证件号", "其他标识", "备注"])
    ws.append(["张三", id1, mobile1, "低保对象"])
    ws.append(["李四", uscc, email, "信访人"])
    ws.append(["王五", bank, plate, "犯罪记录"])
    ws["B2"].comment = Comment(f"曾用证件号：{id3}", "anongw")
    ws["A6"], ws["B6"], ws["C6"] = "赵六", id2, mobile2
    ws.row_dimensions[6].hidden = True                       # 隐藏行
    ws.column_dimensions["D"].hidden = True                  # 隐藏列
    wb.save(path)

    return {
        id1: "Sheet1!B2", mobile1: "Sheet1!C2", "低保对象": "Sheet1!D2",
        uscc: "Sheet1!B3", email: "Sheet1!C3", "信访人": "Sheet1!D3",
        bank: "Sheet1!B4", plate: "Sheet1!C4", "犯罪记录": "Sheet1!D4",
        "赵六": "Sheet1!A6", id2: "Sheet1!B6", mobile2: "Sheet1!C6",
        id3: "Sheet1!B2#comment",
    }


def check_xlsx_selfmade() -> str:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / "selfmade_hidden_comment.xlsx"
    expect = _selfmade_xlsx(path)
    report = inspect_bytes(path.name, path.read_bytes())
    misses = []
    for value, location in expect.items():
        etype = "SENSITIVE_ATTR" if value in ("低保对象", "信访人", "犯罪记录") else (
            "PERSON" if value in ("张三", "李四", "王五", "赵六") else "*")
        hits = [ff for ff in report.findings if ff.location == location
                and (etype == "*" or ff.finding.type.value == etype)
                and (ff.finding.raw == value or ff.finding.normalized == value)]
        if not hits:
            misses.append(f"{location}<-{value[:14]}")
    assert not misses, f"selfmade xlsx misses: {misses}"
    assert report.risk_level == "HIGH", report.risk_level  # 敏感属性 + 批量证件
    names = _matches(report, "PERSON", "张三")
    assert names and names[0].finding.confidence == 1.0  # 词表人名（隐藏面外加分项）
    return f"隐藏列D/隐藏行6/批注 三面 {len(expect)}/{len(expect)} 精确定位命中"


# ── 3. docx 全位面（正文/表格/页眉/页脚）───────────────────────────


def _selfmade_docx(path: Path) -> dict[str, str]:
    """自造 docx：四位面各埋已知值；返回 值→期望定位 表。"""
    import docx

    rng = random.Random(20260930)
    region, _ = nums.pick_region(rng)
    ident = nums.gen_id_card(rng, region, nums.gen_birth(rng))
    mobile = nums.gen_mobile(rng)
    bank, _name = nums.gen_bank_card(rng)
    email = nums.gen_email(rng)
    ipv4 = nums.gen_ipv4(rng)
    landline = "0898-67630001"

    doc = docx.Document()
    doc.add_paragraph("关于张三低保资格核查的联系单")
    doc.add_paragraph(f"核查对象张三（公民身份号码{ident}，联系电话{mobile}），"
                      "请于本周五前反馈。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "户主"
    table.cell(0, 1).text = "银行账户"
    table.cell(1, 0).text = "张三"
    table.cell(1, 1).text = bank
    section = doc.sections[0]
    section.header.paragraphs[0].text = f"澄迈县民政局内部资料 · 联系电话 {landline}"
    section.footer.paragraphs[0].text = f"回执邮箱 {email}，值班室 {ipv4}"
    doc.save(path)

    return {
        ident: "para:2", mobile: "para:2", bank: "table:1:r2:c2",
        "内部资料": "header:1:para:1", landline: "header:1:para:1",
        email: "footer:1:para:1", ipv4: "footer:1:para:1",
    }


def check_docx_selfmade() -> str:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / "selfmade_four_planes.docx"
    expect = _selfmade_docx(path)
    report = inspect_bytes(path.name, path.read_bytes())
    misses = []
    for value, location in expect.items():
        hits = [ff for ff in report.findings if ff.location == location
                and (ff.finding.raw == value or ff.finding.normalized == value)]
        if not hits:
            misses.append(f"{location}<-{value[:14]}")
    assert not misses, f"selfmade docx misses: {misses}"
    kinds = {ff.location.split(":")[0] for ff in report.findings if ff.location}
    assert {"para", "table", "header", "footer"} <= kinds, kinds
    return f"四位面 {len(expect)}/{len(expect)} 精确定位命中（{sorted(kinds)}）"


# ── 4. pdf 坐标级（charbox 并集复核 + 按框取文回读）────────────────


def check_pdf_coords(fixtures: list[dict]) -> str:
    import importlib

    entry = next(e for e in fixtures if e["kind"] == "pdf")
    seeded = next(s for s in entry["seeded"]
                  if _gradeable(s["type"], s["value"]) and s["type"] == "ID_CARD")
    path = DATA_DIR / entry["filename"]
    data = path.read_bytes()
    report = inspect_bytes(path.name, data)

    parsed = parse_pdf(data)
    page_text = parsed.text_layer.page_texts[0]
    value, want = seeded["value"], seeded["value"]
    unions = []
    pos = page_text.find(value)
    while pos != -1:
        box = parsed.text_layer.union(1, pos, pos + len(value))
        assert box is not None
        unions.append(box)
        pos = page_text.find(value, pos + 1)
    assert unions, "seeded value not in extracted text layer"

    hits = _matches(report, "ID_CARD", want)
    assert hits and all(h.bbox is not None for h in hits)
    for h in hits:
        close = [u for u in unions
                 if all(abs(a - b) < 1e-6 for a, b in zip(h.bbox, u, strict=True))]
        assert close, (h.bbox, unions[:2])

    # 库级独立复核：render bbox → 原生系「按框取文」应回读到该值本身
    reader_cfg = yaml.safe_load(
        (REPO_ROOT / "config" / "filechannel.yaml").read_text(encoding="utf-8"))
    pdfium = importlib.import_module(str(reader_cfg["pdf_text_reader"]["module"]))
    doc = pdfium.PdfDocument(data)
    page = doc[0]
    width, height = page.get_size()
    tp = page.get_textpage()
    x0, y0, x1, y1 = hits[0].bbox
    text_in_box = tp.get_text_bounded(left=x0 - 1, bottom=height - y1 - 1,
                                      right=x1 + 1, top=height - y0 + 1)
    doc.close()
    assert text_in_box == value, (text_in_box, value)

    # render 系口径：页内、y0<y1（原点左上）
    assert 0 <= x0 < x1 <= width and 0 <= y0 < y1 <= height, hits[0].bbox
    return (f"{entry['filename']} {value[:12]}… bbox==charbox并集（{len(unions)}处）"
            f"且按框取文回读全等")


def check_scan_pdf() -> str:
    """零文本层 pdf → scan_pdf 如实降级（OCR 路径 D4 接入前零命中）。"""
    import importlib

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / "selfmade_blank.pdf"
    pdfmod = importlib.import_module(str(load_corpus_config()["pdf_writer"]["module"]))
    doc = pdfmod.open()
    doc.new_page(width=595.0, height=842.0)  # 纯空白页：无文本层
    doc.save(str(path), deflate=True, no_new_id=True)
    doc.close()
    report = inspect_bytes(path.name, path.read_bytes())
    assert report.kind == "scan_pdf", report.kind
    assert report.findings == [] and report.risk_level == "NONE"
    assert report.pages == 1
    return "blank pdf → kind=scan_pdf / 零命中 / pages=1"


# ── 5. 三路由材料（与 ops/e2e_smoke SeededMaterial 值层面同源冻结）────
# 说明：本检查不 import ops.e2e_smoke（其模块级 import 牵引网关链路，T3.3 并行
# 开发中），而是持有同一批 seeded 值的独立副本，物化到本 eval 专属目录
# （data/fixtures/filesvc/materials/，不与 e2e 的 data/fixtures/materials/
# 共写同一文件，避免两 eval 互相覆盖）；ops/e2e_smoke 常量如变更需同步本表。


def _valid_id(prefix17: str) -> str:
    """GB11643 校验位补算（与 ops/e2e_smoke._valid_id 同式）。"""
    weights = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
    codes = "10X98765432"
    total = sum(int(ch) * w for ch, w in zip(prefix17, weights, strict=True))
    return prefix17 + codes[total % 11]


_M_ID_A = "11010519491231002X"
_M_ID_C = _valid_id("46002320010307123")
_M_ID_D = _valid_id("11010519491231101")
_M_PHONE = "13800138000"
_M_PERSON = "张三"
_M_PHONE_BLOCKED = "13900139001"

_M_ORDINARY_TEXT = (
    "某镇人民政府办公室\n"
    "关于开展全镇防汛演练的通知\n"
    "\n"
    "各村（居）委会：\n"
    "为检验防汛应急预案的可操作性，定于本周四上午九点在镇政府大院开展防汛演练，\n"
    "请各村网格员提前到场并做好记录，演练结束后将情况汇总报镇党政办。\n"
    "\n"
    "特此通知。"
)
_M_ROSTER_TEXT = (
    "某县民政局低保对象名单（2026年10月·节选）\n"
    "\n"
    "经入户核查与县级联审，现将本月新增低保户名单公示如下：\n"
    f"一、李四，公民身份号码{_M_ID_A}，联系电话{_M_PHONE}；\n"
    f"二、王五，公民身份号码{_M_ID_C}；\n"
    f"三、赵六，公民身份号码{_M_ID_D}。\n"
    "对名单有异议的，请在公示期内向县民政局社会救助科反映。"
)
_M_ROSTER_SPANS: tuple[tuple[str, str, str], ...] = (
    ("李四", "PERSON", "李四"), ("王五", "PERSON", "王五"), ("赵六", "PERSON", "赵六"),
    (_M_ID_A, "ID_CARD", _M_ID_A), (_M_ID_C, "ID_CARD", _M_ID_C),
    (_M_ID_D, "ID_CARD", _M_ID_D), (_M_PHONE, "PHONE_MOBILE", _M_PHONE),
)
_M_CLASSIFIED_TEXT = (
    "机密★某县干部考察纪要（内部资料·注意保密）\n"
    "\n"
    f"考察对象：{_M_PERSON}，联系电话{_M_PHONE_BLOCKED}。\n"
    "考察组意见：该同志政治素质过硬、工作实绩突出，建议进一步培养使用。\n"
    "本纪要不得外传，请按规定归档管理。"
)


def _materialize(name: str, text: str) -> Path:
    """材料物化到本 eval 专属目录（确定性、幂等、回读校验）。"""
    MATERIALS_OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = MATERIALS_OUT_DIR / name
    path.write_text(text, encoding="utf-8", newline="\n")
    assert path.read_text(encoding="utf-8") == text, f"material round-trip: {path}"
    return path


def check_materials() -> str:
    import importlib

    import docx

    notes = []

    # 普通公文 → docx：零命中 / NONE（正文负例控制）
    doc = docx.Document()
    for line in _materialize("ordinary_flood_drill_notice.txt",
                             _M_ORDINARY_TEXT).read_text(encoding="utf-8").splitlines():
        doc.add_paragraph(line)
    path_docx = OUT_DIR / "material_ordinary.docx"
    doc.save(path_docx)
    rep = inspect_bytes(path_docx.name, path_docx.read_bytes())
    assert rep.kind == "docx" and rep.findings == [] and rep.risk_level == "NONE", \
        (len(rep.findings), rep.risk_level)
    notes.append("ordinary→docx 零命中/NONE")

    # 低保名单 → xlsx（名单行进可见列 A，证件/联系方式复制进隐藏列 B）
    from openpyxl import Workbook

    path_xlsx = OUT_DIR / "material_roster.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet1"
    lines = _materialize("low_income_roster_202610.txt",
                         _M_ROSTER_TEXT).read_text(encoding="utf-8").splitlines()
    for i, line in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=line)
    hidden_vals = [s for s in _M_ROSTER_SPANS if s[1] in ("ID_CARD", "PHONE_MOBILE")]
    for i, (_raw, _etype, normalized) in enumerate(hidden_vals, 1):
        ws.cell(row=i, column=2, value=normalized)
    ws.column_dimensions["B"].hidden = True
    wb.save(path_xlsx)
    rep = inspect_bytes(path_xlsx.name, path_xlsx.read_bytes())
    misses = []
    for raw, etype, normalized in _M_ROSTER_SPANS:
        if not _matches(rep, etype, normalized):
            misses.append(f"{etype}:{raw[:12]}")
    assert not misses, f"roster span misses: {misses}"
    for _raw, etype, normalized in hidden_vals:
        at_hidden = [ff for ff in _matches(rep, etype, normalized)
                     if ff.location and re.fullmatch(r"Sheet1!B\d+", ff.location)]
        assert at_hidden, f"hidden col miss: {normalized[:12]}"
    assert rep.risk_level == "HIGH", rep.risk_level
    notes.append(f"roster→xlsx spans {len(_M_ROSTER_SPANS)}"
                 f"/{len(_M_ROSTER_SPANS)} + 隐藏列B/HIGH")

    # 机密纪要 → pdf 文本层：密级词 + 人名/手机号全命中、bbox 齐备
    path_pdf = OUT_DIR / "material_classified.pdf"
    pdfmod = importlib.import_module(
        str(load_corpus_config()["pdf_writer"]["module"]))
    doc_pdf = pdfmod.open()
    page = doc_pdf.new_page(width=595.0, height=842.0)
    y = 72.0
    for line in _materialize("classified_cadre_minutes.txt",
                             _M_CLASSIFIED_TEXT).read_text(encoding="utf-8").splitlines():
        if line:
            page.insert_text((72.0, y), line, fontname="china-s", fontsize=10)
        y += 16.0
    doc_pdf.save(str(path_pdf), deflate=True, no_new_id=True)
    doc_pdf.close()
    rep = inspect_bytes(path_pdf.name, path_pdf.read_bytes())
    assert rep.kind == "pdf" and rep.pages == 1
    marks = {ff.finding.raw for ff in rep.findings
             if ff.finding.type is EntityClass.CLASSIFICATION_MARK}
    assert {"机密★", "内部资料", "注意保密", "不得外传"} <= marks, marks
    assert _matches(rep, "PERSON", _M_PERSON), "机密材料人名未命中"
    assert _matches(rep, "PHONE_MOBILE", _M_PHONE_BLOCKED), "机密材料手机号未命中"
    assert all(ff.bbox is not None and ff.page == 1 for ff in rep.findings)
    # §5.5 字面口径：密级词不进分级阶梯，MID 由结构化 PII（人名/手机号）决定
    assert rep.risk_level == "MID", rep.risk_level
    notes.append(f"classified→pdf 密级词{len(marks)}类+bbox/MID")
    return "；".join(notes)


# ── 6. 风险分级（§5.5 纯函数逐档）──────────────────────────────────


def _ff(etype: EntityClass, *, whitelisted: bool = False) -> FileFinding:
    return FileFinding(page=1, location="para:1", finding=Finding(
        fid="f_0001", type=etype, raw="x", normalized="x", whitelisted=whitelisted))


def check_risk_ladder() -> str:
    cases = [
        ([], "NONE"),
        ([_ff(EntityClass.DOC_NUMBER, whitelisted=True)], "LOW"),
        ([_ff(EntityClass.CLASSIFICATION_MARK)], "LOW"),       # 密级词不进阶梯
        ([_ff(EntityClass.PERSON)], "MID"),
        ([_ff(EntityClass.PHONE_MOBILE)], "MID"),
        ([_ff(EntityClass.SENSITIVE_ATTR)], "HIGH"),
        ([_ff(EntityClass.ID_CARD) for _ in range(3)], "HIGH"),  # 批量线=3
        ([_ff(EntityClass.BANK_CARD) for _ in range(3)], "HIGH"),
        ([_ff(EntityClass.ID_CARD), _ff(EntityClass.ID_CARD)], "MID"),
        ([_ff(EntityClass.ID_CARD, whitelisted=True) for _ in range(5)], "LOW"),
    ]
    for findings, want in cases:
        got = risk_level_of(findings, batch_threshold=3)
        assert got == want, ([f.finding.type.value for f in findings], got, want)
    return f"{len(cases)} 档全对（HIGH/MID/LOW/NONE + 白名单豁免 + 批量线）"


# ── 7. 边界（大小/种类/畸形）───────────────────────────────────────


def check_limits() -> str:
    try:
        inspect_bytes("big.pdf", b"\0" * (th.FILE_SIZE_MAX_BYTES + 1))
        raise AssertionError("50MB+1 未拒绝")
    except FileTooLargeError:
        pass
    for name in ("x.doc", "noext", "x.xlsx.txt"):
        try:
            inspect_bytes(name, b"whatever")
            raise AssertionError(f"{name} 未拒绝")
        except UnsupportedFileType:
            pass
    for name, blob in (("broken.docx", b"PK\x03\x04broken"),
                       ("broken.xlsx", b"PK\x03\x04broken"),
                       ("broken.pdf", b"%PDF-1.7 broken")):
        try:
            inspect_bytes(name, blob)
            raise AssertionError(f"{name} 未报解析错误")
        except (DocumentParseError, UnsupportedFileType):
            pass
    return "50MB 上限 / 不支持种类 / 畸形文件 全部异常语义拒绝"


# ── 入口 ───────────────────────────────────────────────────────────


def main() -> int:
    fixtures = _load_fixtures()
    _record("fixtures(15份/sha/kind/fid/命中率100%)", lambda: check_fixtures(fixtures))
    _record("xlsx-hidden(M8夹具隐藏列D)", lambda: check_xlsx_hidden_col(fixtures))
    _record("xlsx-selfmade(隐藏列/隐藏行/批注精确定位)", check_xlsx_selfmade)
    _record("docx-selfmade(正文/表格/页眉/页脚四位面)", check_docx_selfmade)
    _record("pdf-coords(charbox并集复核+按框取文)", lambda: check_pdf_coords(fixtures))
    _record("scan-pdf(零文本层如实降级)", check_scan_pdf)
    _record("materials(三路由材料docx/xlsx/pdf命中)", check_materials)
    _record("risk-ladder(§5.5 逐档纯函数)", check_risk_ladder)
    _record("limits(50MB/种类/畸形)", check_limits)

    for _, line in RESULTS:
        print(line)
    passed = sum(1 for ok, _ in RESULTS if ok)
    total = len(RESULTS)
    print(f"M6 FILESVC(text-layer): {passed}/{total} checks passed")
    return 0 if passed == total else 1


if __name__ == "__main__":
    sys.exit(main())
